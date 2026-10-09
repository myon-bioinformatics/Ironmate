from pathlib import Path

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from pptx_replace import transform


def fixture(tmp_path):
    path = tmp_path / "source.pptx"
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = "Draft title"
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.font.bold = True
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    p.save(path)
    return path


def test_dry_run_and_preservation(tmp_path):
    src = fixture(tmp_path)
    dst = tmp_path / "output.pptx"
    rules = [{"operation": "replace_text", "old": "Draft", "new": "Final"}, {"operation": "text_color", "color": "#000000"}]
    preview = transform(src, rules, dry_run=True)
    assert preview["change_count"] == 2
    assert not dst.exists()
    result = transform(src, rules, destination=dst)
    assert result["change_count"] == 2
    original = Presentation(src).slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    updated = Presentation(dst).slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert original.text == "Draft title"
    assert updated.text == "Final title"
    assert updated.font.bold is True
    assert updated.font.color.rgb == RGBColor(0, 0, 0)


def test_background_and_no_overwrite(tmp_path):
    src = fixture(tmp_path)
    dst = tmp_path / "out.pptx"
    assert transform(src, [{"operation": "slide_background", "color": "#112233"}], destination=dst)["change_count"] == 1
    assert Presentation(dst).slides[0].background.fill.fore_color.rgb == RGBColor(17, 34, 51)
    with pytest.raises(ValueError, match="new file"):
        transform(src, [{"operation": "replace_text", "old": "none", "new": "x"}], destination=dst)


def test_unsupported_and_zero_match(tmp_path):
    src = fixture(tmp_path)
    assert transform(src, [{"operation": "replace_text", "old": "missing", "new": "x"}], dry_run=True)["change_count"] == 0
    with pytest.raises(ValueError, match="unsupported operation"):
        transform(src, [{"operation": "arrow", "color": "#000000"}], dry_run=True)


def test_cli_one_liner_and_explicit_failure_evidence(tmp_path):
    """An intentionally failing child suite must yield a nonzero exit and JUnit failure."""
    import subprocess
    import sys
    import xml.etree.ElementTree as ET

    src = fixture(tmp_path)
    output = tmp_path / "cli-output.pptx"
    cli = Path(__file__).resolve().parents[1] / "pptx_replace.py"
    command = [sys.executable, str(cli), str(src), "--replace", "Draft", "Final", "--output", str(output)]
    completed = subprocess.run(command, capture_output=True, text=True, check=False)
    assert completed.returncode == 0, completed.stderr
    assert Presentation(output).slides[0].shapes[0].text_frame.paragraphs[0].runs[0].text == "Final title"

    child = tmp_path / "test_intentional_failure.py"
    child.write_text(
        "def test_expected_failure():\\n    assert 1 == 2, 'intentional pptx CI failure probe'\\n".replace("\\n", "\n"),
        encoding="utf-8",
    )
    junit = tmp_path / "intentional-failure.xml"
    probe = subprocess.run(
        [sys.executable, "-m", "pytest", "-q", str(child), f"--junitxml={junit}"],
        capture_output=True, text=True, check=False, cwd=tmp_path,
    )
    assert probe.returncode == 1, probe.stdout + probe.stderr
    root = ET.parse(junit).getroot()
    failures = root.findall(".//testcase/failure")
    assert len(failures) == 1
    assert "intentional pptx CI failure probe" in (failures[0].text or "")


def test_output_collision_is_rejected_before_reading_input(tmp_path):
    existing = tmp_path / "existing.pptx"
    existing.write_bytes(b"preserve me")
    with pytest.raises(ValueError, match="new file"):
        transform(tmp_path / "missing.pptx", [{"operation": "replace_text", "old": "x", "new": "y"}], destination=existing)
    assert existing.read_bytes() == b"preserve me"


def test_slide_xml_preserves_unrelated_shape_and_text(tmp_path):
    import zipfile
    import xml.etree.ElementTree as ET

    src = fixture(tmp_path)
    dst = tmp_path / "changed.pptx"
    transform(src, [{"operation": "replace_text", "old": "Draft", "new": "Final"}], destination=dst)
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    def slide_xml(path):
        with zipfile.ZipFile(path) as archive:
            return ET.fromstring(archive.read("ppt/slides/slide1.xml"))
    before, after = slide_xml(src), slide_xml(dst)
    assert [node.text for node in before.findall(".//a:t", ns)] == ["Draft title"]
    assert [node.text for node in after.findall(".//a:t", ns)] == ["Final title"]
    # Normalized XML trees must agree after restoring only the intended text change.
    after.find(".//a:t", ns).text = "Draft title"
    assert ET.tostring(before) == ET.tostring(after)


def test_font_size_bold_and_rgb_xml_roundtrip(tmp_path):
    import zipfile
    import xml.etree.ElementTree as ET

    src = fixture(tmp_path)
    dst = tmp_path / "styled.pptx"
    rules = [
        {"operation": "font_size", "points": 24},
        {"operation": "bold", "value": False},
        {"operation": "text_color", "color": "#112233"},
    ]
    assert transform(src, rules, dry_run=True)["change_count"] == 3
    assert transform(src, rules, destination=dst)["change_count"] == 3
    run = Presentation(dst).slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.size.pt == 24
    assert run.font.bold is False
    assert run.font.color.rgb == RGBColor(17, 34, 51)
    with zipfile.ZipFile(dst) as z:
        root = ET.fromstring(z.read("ppt/slides/slide1.xml"))
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    properties = root.find(".//a:rPr", ns)
    assert properties is not None
    assert properties.attrib["sz"] == "2400"
    assert properties.attrib["b"] == "0"
    assert properties.find(".//a:srgbClr", ns).attrib["val"].upper() == "112233"


@pytest.mark.parametrize("rule", [
    {"operation": "font_size", "points": 0},
    {"operation": "font_size", "points": "12"},
    {"operation": "font_size", "points": True},
    {"operation": "bold", "value": "true"},
])
def test_style_rule_rejects_invalid_types(tmp_path, rule):
    src = fixture(tmp_path)
    with pytest.raises(ValueError):
        transform(src, [rule], dry_run=True)


def test_batch_font_family_override(tmp_path):
    src = fixture(tmp_path)
    dst = tmp_path / "font.pptx"
    rules = [{"operation": "font_family", "name": "Yu Gothic"}]
    assert transform(src, rules, dry_run=True)["change_count"] == 1
    assert transform(src, rules, destination=dst)["change_count"] == 1
    run = Presentation(dst).slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.name == "Yu Gothic"
    import zipfile
    import xml.etree.ElementTree as ET
    with zipfile.ZipFile(dst) as archive:
        root = ET.fromstring(archive.read("ppt/slides/slide1.xml"))
    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    latin = root.find(".//a:rPr/a:latin", ns)
    assert latin is not None and latin.attrib["typeface"] == "Yu Gothic"


def test_font_family_rejects_empty(tmp_path):
    src = fixture(tmp_path)
    with pytest.raises(ValueError, match="nonempty"):
        transform(src, [{"operation": "font_family", "name": ""}], dry_run=True)
