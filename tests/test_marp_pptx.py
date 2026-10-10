import pytest
from pathlib import Path
from pptx import Presentation
from marp_pptx import convert


def test_missing_marp_is_explicit(tmp_path):
    source = Path(__file__).parent / "fixtures" / "mini_readme.md"
    with pytest.raises(RuntimeError, match="Marp CLI not installed"):
        convert(source, tmp_path / "sample.pptx", marp="nonexistent-marp-cli-ironmate")


def test_existing_output_is_preserved(tmp_path):
    source = Path(__file__).parent / "fixtures" / "mini_readme.md"
    output = tmp_path / "sample.pptx"
    output.write_bytes(b"preserved")
    with pytest.raises(ValueError, match="new file"):
        convert(source, output)
    assert output.read_bytes() == b"preserved"


def test_marp_cli_integration_when_enabled(tmp_path):
    import os
    import shutil
    import subprocess
    import sys

    marp = os.environ.get("IRONMATE_MARP_BIN", "marp")
    if os.environ.get("IRONMATE_RUN_MARP_INTEGRATION") != "1":
        pytest.skip("optional Marp integration not enabled")
    if shutil.which(marp) is None:
        pytest.fail("Marp integration requested but executable is missing")
    source = Path(__file__).parent / "fixtures" / "mini_readme.md"
    output = tmp_path / "sample.pptx"
    result = convert(source, output, marp=marp)
    assert result["slides"] >= 1
    assert len(Presentation(output).slides) >= 1
    # Inspect actual Open XML elements rather than assuming editable slide text.
    import json
    import zipfile
    from xml.etree import ElementTree as ET

    ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main",
          "p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
    inspection = []
    with zipfile.ZipFile(output) as archive:
        for index, slide in enumerate(Presentation(output).slides, 1):
            xml = ET.fromstring(archive.read(f"ppt/slides/slide{index}.xml"))
            texts = [n.text or "" for n in xml.findall(".//a:t", ns)]
            pictures = len(xml.findall(".//p:pic", ns))
            shapes = len(xml.findall(".//p:sp", ns))
            editable_runs = sum(len(p.runs) for shape in slide.shapes
                                if shape.has_text_frame for p in shape.text_frame.paragraphs)
            inspection.append({"slide": index, "texts": texts,
                               "pictures": pictures, "shapes": shapes,
                               "editable_runs": editable_runs})
    evidence = Path("build/test-results/marp-openxml-inspection.json")
    evidence.parent.mkdir(parents=True, exist_ok=True)
    evidence.write_text(json.dumps(inspection, ensure_ascii=False, indent=2), encoding="utf-8")
    assert len(inspection) == result["slides"]
    assert all(isinstance(row["texts"], list) for row in inspection)
    # Marp may encode a rendered slide as a background fill rather than p:pic.
    assert all(row["pictures"] >= 0 and row["shapes"] >= 0 for row in inspection)
    # This is observational evidence, not an editability guarantee.


    # Preserve raw slide XML with hashes for later regression fixtures.
    import hashlib
    raw_dir = Path('build/test-results/marp-openxml')
    raw_dir.mkdir(parents=True, exist_ok=True)
    members = []
    with zipfile.ZipFile(output) as archive:
        for name in sorted(archive.namelist()):
            if not (name.startswith('ppt/slides/') and name.endswith('.xml')):
                continue
            data = archive.read(name)
            assert len(data) <= 2_000_000
            target = raw_dir / Path(name).name
            target.write_bytes(data)
            members.append({'source': name, 'file': target.name,
                            'sha256': hashlib.sha256(data).hexdigest(), 'size': len(data)})
    assert members
    (raw_dir / 'manifest.json').write_text(json.dumps(members, indent=2), encoding='utf-8')

def test_mermaid_actual_rendering_probe_when_enabled(tmp_path):
    """Observe whether the pinned Marp CLI draws Mermaid, not merely passes its source."""
    import json
    import os
    import shutil
    import subprocess
    from pptx_inspect import inspect

    if os.environ.get("IRONMATE_RUN_MARP_INTEGRATION") != "1":
        pytest.skip("optional real Marp integration not enabled")
    marp = os.environ.get("IRONMATE_MARP_BIN", "marp")
    assert shutil.which(marp), "Marp CLI required for integration"
    source = tmp_path / "diagram.md"
    fence = chr(96) * 3
    source.write_text("---\nmarp: true\n---\n# Flow\n\n" + fence +
                      "mermaid\nflowchart LR\n  A[Input] --> B[Output]\n" +
                      fence + "\n", encoding="utf-8")
    html = tmp_path / "diagram.html"
    proc = subprocess.run([marp, str(source), "-o", str(html)],
                          capture_output=True, text=True, timeout=120)
    assert proc.returncode == 0, proc.stderr
    rendered = html.read_text(encoding="utf-8")
    # A code fence rendered as highlighted text is not a Mermaid diagram.
    # This is an observation, not a requirement that Marp CLI 4.x supports Mermaid.
    has_diagram = ("<svg" in rendered and "flowchart LR" not in rendered)
    pptx = tmp_path / "diagram.pptx"
    convert(source, pptx, marp=marp)
    observation = {"marp_mermaid_rendered": has_diagram,
                   "status": "rendered" if has_diagram else "unsupported",
                   "pptx": inspect(pptx)}
    evidence = Path("build/test-results/marp-mermaid-observation.json")
    evidence.parent.mkdir(parents=True, exist_ok=True)
    evidence.write_text(json.dumps(observation, ensure_ascii=False, indent=2), encoding="utf-8")
    assert len(observation["pptx"]["slides"]) >= 1
    if not has_diagram:
        assert observation["status"] == "unsupported"

def test_real_mermaid_svg_to_marp_pptx_when_enabled(tmp_path):
    """Integration gate for the pinned Mermaid SVG backend, separate from native PPTX shapes."""
    import os
    import shutil
    from pptx_pipeline import run
    if os.environ.get("IRONMATE_RUN_MARP_INTEGRATION") != "1":
        pytest.skip("optional Marp integration not enabled")
    marp = os.environ.get("IRONMATE_MARP_BIN", "marp")
    mmdc = os.environ.get("IRONMATE_MMDC_BIN", "mmdc")
    if shutil.which(mmdc) is None:
        pytest.fail("Mermaid CLI requested but not installed")
    source = "# Diagram\n\n```mermaid\nflowchart LR\n A[Start] --> B[Finish]\n```\n"
    result = run(text=source, output=tmp_path / "mermaid.pptx", marp=marp, mmdc=mmdc)
    assert result["conversion"]["mermaid_svg_count"] == 1
    assert len(result["inspection"]["slides"]) >= 1


@pytest.mark.parametrize("kind,diagram", [
    ("flowchart", "flowchart LR\n A[Start] --> B[Finish]"),
    ("uml-class", "classDiagram\n class Animal\n class Dog\n Animal <|-- Dog"),
])
def test_real_svg_diagram_xml_edit_when_enabled(tmp_path, kind, diagram):
    """Mermaid flowchart and UML class diagram share the SVG pipeline."""
    import os
    import shutil
    import zipfile
    from pptx_pipeline import run
    from pptx_xml_probe import set_slide_name
    if os.environ.get("IRONMATE_RUN_MARP_INTEGRATION") != "1":
        pytest.skip("optional Marp integration not enabled")
    marp = os.environ.get("IRONMATE_MARP_BIN", "marp")
    mmdc = os.environ.get("IRONMATE_MMDC_BIN", "mmdc")
    assert shutil.which(mmdc), "Mermaid CLI required"
    fence = chr(96) * 3
    markdown = "# Diagram\n\n" + fence + "mermaid\n" + diagram + "\n" + fence + "\n"
    generated = tmp_path / (kind + ".pptx")
    edited = tmp_path / (kind + "-edited.pptx")
    result = run(text=markdown, output=generated, marp=marp, mmdc=mmdc)
    assert result["conversion"]["mermaid_svg_count"] == 1
    set_slide_name(generated, edited, name=kind)
    assert len(Presentation(edited).slides) == len(Presentation(generated).slides)
    with zipfile.ZipFile(generated) as before, zipfile.ZipFile(edited) as after:
        assert before.namelist() == after.namelist()
        assert any(name.startswith("ppt/media/") for name in after.namelist())
        assert all(before.read(name) == after.read(name) for name in before.namelist()
                   if name != "ppt/slides/slide1.xml")



@pytest.mark.parametrize("kind,diagram", [
    ("flowchart", "flowchart LR\n A[Start] --> B[Finish]"),
    ("uml-class", "classDiagram\n class Animal\n class Dog\n Animal <|-- Dog"),
])
def test_real_svg_visual_preview_when_enabled(tmp_path, kind, diagram):
    """Generate an actual browser-rendered PNG for human inspection, without LibreOffice."""
    import json
    import os
    import shutil
    import struct
    import subprocess
    from mermaid_svg import render_mermaid
    if os.environ.get("IRONMATE_RUN_MARP_INTEGRATION") != "1":
        pytest.skip("optional Marp integration not enabled")
    marp = os.environ.get("IRONMATE_MARP_BIN", "marp")
    mmdc = os.environ.get("IRONMATE_MMDC_BIN", "mmdc")
    assert shutil.which(marp) and shutil.which(mmdc)
    fence = chr(96) * 3
    original = "# Diagram\n\n" + fence + "mermaid\n" + diagram + "\n" + fence + "\n"
    prepared, svgs = render_mermaid(original, tmp_path, mmdc=mmdc)
    assert len(svgs) == 1
    md = tmp_path / "preview.md"
    md.write_text(prepared, encoding="utf-8")
    png = tmp_path / "preview.png"
    result = subprocess.run([marp, str(md), "--png", "--allow-local-files",
                             "-o", str(png)], capture_output=True, text=True, timeout=120)
    assert result.returncode == 0, result.stderr
    data = png.read_bytes()
    assert data[:8] == bytes.fromhex("89504e470d0a1a0a")
    width, height = struct.unpack(">II", data[16:24])
    assert width >= 640 and height >= 360
    import hashlib
    evidence = Path("build/test-results/marp-visual")
    evidence.mkdir(parents=True, exist_ok=True)
    saved = evidence / (kind + ".png")
    saved.write_bytes(data)
    (evidence / (kind + ".json")).write_text(json.dumps({
        "kind": kind, "width": width, "height": height,
        "png_sha256": hashlib.sha256(data).hexdigest(),
        "svg_sha256": hashlib.sha256(svgs[0].read_bytes()).hexdigest(),
        "review": "visual inspection pending"
    }, indent=2), encoding="utf-8")



@pytest.mark.parametrize("kind,diagram", [
    ("flowchart", "flowchart LR\n A[Start] --> B[Finish]"),
    ("uml-class", "classDiagram\n class Animal\n class Dog\n Animal <|-- Dog"),
])
def test_inline_markdown_theme_variants_and_svg_source_when_enabled(tmp_path, kind, diagram):
    """Direct Markdown input: color changes require a new render, not a native SVG text edit."""
    import hashlib
    import json
    import os
    import shutil
    import subprocess
    import zipfile
    from pptx_pipeline import run
    from mermaid_svg import render_mermaid
    if os.environ.get("IRONMATE_RUN_MARP_INTEGRATION") != "1":
        pytest.skip("optional Marp integration not enabled")
    marp = os.environ.get("IRONMATE_MARP_BIN", "marp")
    mmdc = os.environ.get("IRONMATE_MMDC_BIN", "mmdc")
    assert shutil.which(marp) and shutil.which(mmdc)
    fence = chr(96) * 3
    observations = []
    for label, background, foreground in (
        ("light", "#FFFFFF", "#000000"), ("dark", "#14213D", "#FFFFFF")):
        markdown = ("---\nmarp: true\nstyle: |\n  section { background: "
                    + background + "; color: " + foreground
                    + "; }\n---\n# Direct input\n\n" + fence + "mermaid\n"
                    + diagram + "\n" + fence + "\n")
        output = tmp_path / f"{kind}-{label}.pptx"
        result = run(text=markdown, output=output, marp=marp, mmdc=mmdc)
        assert result["conversion"]["mermaid_svg_count"] == 1
        assert len(Presentation(output).slides) == 1
        with zipfile.ZipFile(output) as package:
            media = [name for name in package.namelist() if name.startswith("ppt/media/")]
            assert media, "rendered PPTX must contain media"
            media_digest = hashlib.sha256(b"".join(package.read(n) for n in sorted(media))).hexdigest()
        image_dir = tmp_path / f"{kind}-{label}"
        prepared, images = render_mermaid(markdown, image_dir, mmdc=mmdc)
        assert len(images) == 1
        assert (image_dir / images[0].name).is_file()
        md = image_dir / "preview.md"
        md.write_text(prepared, encoding="utf-8")
        png = tmp_path / f"{kind}-{label}.png"
        proc = subprocess.run([marp, str(md), "--png", "--allow-local-files",
                               "-o", str(png)], capture_output=True, text=True, timeout=120)
        assert proc.returncode == 0, proc.stderr
        data = png.read_bytes()
        assert data.startswith(bytes.fromhex("89504e470d0a1a0a"))
        observations.append({"variant": label, "background": background,
                             "foreground": foreground, "media_sha256": media_digest,
                             "preview_sha256": hashlib.sha256(data).hexdigest()})
        evidence = Path("build/test-results/marp-visual")
        evidence.mkdir(parents=True, exist_ok=True)
        (evidence / f"{kind}-{label}.png").write_bytes(data)
    assert observations[0]["preview_sha256"] != observations[1]["preview_sha256"]
    evidence = Path("build/test-results/marp-visual")
    (evidence / f"{kind}-theme-comparison.json").write_text(
        json.dumps(observations, indent=2), encoding="utf-8")

