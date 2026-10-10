"""Open XML shape mutation preserves other package members and rejects ambiguity."""
from zipfile import ZipFile
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches
import pytest
from pptx_xml import edit_shape


def sample(path, *, explicit=True):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    shape.name = "Target"
    if explicit:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(20, 30, 40)
        shape.line.color.rgb = RGBColor(50, 60, 70)
    prs.save(path)


def test_shape_fill_line_and_geometry_roundtrip(tmp_path):
    src, dst = tmp_path / "source.pptx", tmp_path / "edited.pptx"
    sample(src)
    edit_shape(src, dst, shape_name="Target", fill_rgb="#112233",
               line_rgb="#445566", left=Inches(2), top=Inches(3),
               width=Inches(4), height=Inches(2))
    shape = Presentation(dst).slides[0].shapes[0]
    assert shape.fill.fore_color.rgb == RGBColor(17, 34, 51)
    assert shape.line.color.rgb == RGBColor(68, 85, 102)
    assert (shape.left, shape.top, shape.width, shape.height) == (
        Inches(2), Inches(3), Inches(4), Inches(2))
    with ZipFile(src) as before, ZipFile(dst) as after:
        assert before.namelist() == after.namelist()
        assert all(before.read(name) == after.read(name) for name in before.namelist()
                   if name != "ppt/slides/slide1.xml")


def test_missing_explicit_style_fails_closed(tmp_path):
    src, dst = tmp_path / "source.pptx", tmp_path / "edited.pptx"
    sample(src, explicit=False)
    with pytest.raises(ValueError, match="explicit RGB"):
        edit_shape(src, dst, shape_name="Target", fill_rgb="#FFFFFF")
    assert not dst.exists()


@pytest.mark.parametrize("kwargs", [
    {"shape_name": "Missing", "fill_rgb": "#112233"},
    {"shape_name": "Target", "width": 0},
    {"shape_name": "Target", "line_rgb": "red"},
    {"shape_name": "Target"},
])
def test_invalid_shape_rules_are_rejected(tmp_path, kwargs):
    src, dst = tmp_path / "source.pptx", tmp_path / "edited.pptx"
    sample(src)
    with pytest.raises(ValueError):
        edit_shape(src, dst, **kwargs)
    assert not dst.exists()

