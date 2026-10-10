"""Regression tests for conservative editability decisions."""
from pptx import Presentation
from pptx.util import Inches
from pptx_inspect import inspect


def test_native_text_is_editable(tmp_path):
    path = tmp_path / "native.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(1)).text = "Hello"
    prs.save(path)
    result = inspect(path)["slides"][0]
    assert result["status"] == "editable"
    assert result["editable_runs"] == 1


def test_no_text_is_unsupported(tmp_path):
    path = tmp_path / "empty.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(path)
    result = inspect(path)["slides"][0]
    assert result["status"] == "unsupported"
    assert result["editable_runs"] == 0
