"""Geometry-only contracts for explicit edge-to-edge connections.

These tests deliberately separate mathematical orientation from PPTX shape
rotation, which must be calibrated using rendered output.
"""
import math
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from tests.test_pptx_diagram_shapes import edge_center

EDGES = ("left", "right", "top", "bottom")


def direction(source, source_edge, target, target_edge):
    sx, sy = edge_center(source, source_edge)
    tx, ty = edge_center(target, target_edge)
    dx, dy = tx - sx, ty - sy
    if dx == 0 and dy == 0:
        raise ValueError("coincident anchors")
    return (sx, sy), (tx, ty), math.degrees(math.atan2(dy, dx))


def head_vertices(center, angle, length, half_width):
    """Return tip, left base, right base in slide coordinate space."""
    theta = math.radians(angle)
    ux, uy = math.cos(theta), math.sin(theta)
    nx, ny = -uy, ux
    cx, cy = center
    tip = (cx + ux * length / 2, cy + uy * length / 2)
    base = (cx - ux * length / 2, cy - uy * length / 2)
    return tip, (base[0] + nx * half_width, base[1] + ny * half_width), (
        base[0] - nx * half_width, base[1] - ny * half_width)


@pytest.mark.parametrize("start_edge", EDGES)
@pytest.mark.parametrize("end_edge", EDGES)
@pytest.mark.parametrize("reverse", (False, True))
def test_all_edge_pairs_and_reverse(start_edge, end_edge, reverse):
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[6])
    a = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    b = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(6), Inches(4), Inches(2), Inches(1))
    source, target = (b, a) if reverse else (a, b)
    start, end, angle = direction(source, start_edge, target, end_edge)
    center = ((start[0] + end[0]) / 2, (start[1] + end[1]) / 2)
    tip, left, right = head_vertices(center, angle, Inches(0.4), Inches(0.15))
    ux, uy = math.cos(math.radians(angle)), math.sin(math.radians(angle))
    assert math.isclose((tip[0] - center[0]) * ux + (tip[1] - center[1]) * uy,
                        Inches(0.2), abs_tol=1)
    assert math.isclose((left[0] + right[0]) / 2, center[0] - ux * Inches(0.2), abs_tol=1)
    assert math.isclose((left[1] + right[1]) / 2, center[1] - uy * Inches(0.2), abs_tol=1)
    assert math.isclose(math.hypot(left[0] - center[0], left[1] - center[1]),
                        math.hypot(right[0] - center[0], right[1] - center[1]), abs_tol=1)


def test_coincident_anchor_rejected():
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[6])
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    with pytest.raises(ValueError, match="coincident"):
        direction(shape, "left", shape, "left")
