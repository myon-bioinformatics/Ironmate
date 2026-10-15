"""Real HTML/PDF preview integration, independent of any external viewer."""
import os
from pathlib import Path

from pptx import Presentation
import pytest

import pptx_preview


@pytest.mark.parametrize('patch', [dict(expected_slides=0), dict(expected_slides=True), dict(expected_slides=101), dict(sha='short'), dict(run_id='')])
def test_preview_preflight_rejects_bad_identity_and_counts(tmp_path, patch):
    html = tmp_path / 'preview.html'
    html.write_text('<html></html>')
    config = dict(expected_slides=1, sha='a' * 40, run_id='test')
    config.update(patch)
    with pytest.raises(ValueError):
        pptx_preview.render_preview_html(html, tmp_path, **config)
    assert not (tmp_path / 'slides').exists()


@pytest.mark.skipif(os.environ.get('IRONMATE_RUN_REVIEW') != '1',
                    reason='optional real Chromium/Marp review lane not enabled')
def test_real_marp_review_html_pdf_and_final_native_table(tmp_path):
    import hashlib
    from pypdf import PdfReader
    import json
    import subprocess
    import sys
    source = tmp_path / 'slides.md'
    source.write_text('---\nmarp: true\n---\n# Alpha\n\n| Key | Value |\n| --- | --- |\n| Source | 10 |\n\n---\n\n# Beta\n\nSecond source slide.\n')
    before = source.read_bytes()
    out = Path(os.environ.get('IRONMATE_REVIEW_EVIDENCE', str(tmp_path))) / 'marp-review'
    out.parent.mkdir(parents=True, exist_ok=True)
    table_source = tmp_path / 'native.csv'
    table_source.write_text('Native header\nPPTX_NATIVE_ONLY\n')
    cli = Path(__file__).resolve().parents[1] / 'pptx_pipeline.py'
    process = subprocess.run([sys.executable, str(cli), str(source), '--review', str(out),
                              '--marp', os.environ['IRONMATE_MARP_BIN'],
                              '--review-sha', os.environ.get('IRONMATE_TEST_SHA', 'a' * 40),
                              '--review-run-id', os.environ.get('GITHUB_RUN_ID', 'local'),
                              '--table', str(table_source), '--table-slide', '2',
                              '--table-region', '1', '4', '8', '2'],
                             capture_output=True, text=True, check=True, timeout=180)
    report = json.loads(process.stdout)
    assert source.read_bytes() == before
    assert report['stage'] == 'complete'
    preview = report['preview']
    assert preview['project'] == 'chromium-marp-preview'
    assert preview['slides'] == preview['pdf_pages'] == 2
    assert not preview['is_final_pptx_render'] and not preview['native_pptx_edits_reflected']
    pdf = PdfReader(out / 'preview.pdf', strict=True)
    assert len(pdf.pages) == 2
    assert 'Alpha' in pdf.pages[0].extract_text()
    assert 'Beta' in pdf.pages[1].extract_text()
    assert 'PPTX_NATIVE_ONLY' not in ''.join(p.extract_text() for p in pdf.pages)
    table = Presentation(out / 'final.pptx').slides[1].shapes[-1].table
    assert table.cell(1, 0).text == 'PPTX_NATIVE_ONLY'
    assert preview['captures'][0]['sha256'] != preview['captures'][1]['sha256']
    for entry in preview['captures']:
        data = (out / entry['screenshot']).read_bytes()
        assert hashlib.sha256(data).hexdigest() == entry['sha256']
        import struct
        w, h = struct.unpack('>II', data[16:24])
        assert w >= 640 and h >= 360 and abs(w / h - 16 / 9) < .01


@pytest.mark.skipif(os.environ.get('IRONMATE_RUN_REVIEW') != '1',
                    reason='optional real Chromium preview lane not enabled')
@pytest.mark.parametrize('problem', ['broken-image', 'wrong-count', 'bad-pdf'])
def test_real_preview_rejects_bad_assets_page_counts_and_pdf(tmp_path, monkeypatch, problem):
    # Test only the HTML preview adapter, without assuming a PowerPoint viewer.
    directory = tmp_path / 'preview'
    directory.mkdir()
    html = directory / 'preview.html'
    image = '<img src="missing.png">' if problem == 'broken-image' else ''
    html.write_text('<html><head></head><body>' + image +
                    '<svg data-marpit-svg="" viewBox="0 0 800 450" width="800" height="450"><rect width="800" height="450" fill="white"/></svg></body></html>')
    if problem == 'bad-pdf':
        from playwright.sync_api import Page
        monkeypatch.setattr(Page, 'pdf', lambda self, **kw: Path(kw['path']).write_bytes(b'%PDF-broken'))
    from pypdf.errors import PdfReadError
    expected = PdfReadError if problem == 'bad-pdf' else ValueError
    match = {'broken-image': 'broken image', 'wrong-count': 'slide counts', 'bad-pdf': None}[problem]
    with pytest.raises(expected, match=match):
        pptx_preview.render_preview_html(html, directory, expected_slides=2 if problem == 'wrong-count' else 1,
                                    sha='a' * 40, run_id='reject', timeout=1500)
    assert not (directory / 'review.json').exists()  # Only review() may publish a completion record.
