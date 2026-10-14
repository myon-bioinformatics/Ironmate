"""Append editable card, region, agenda, flow and table patterns to PPTX slides."""
import argparse
import csv
import json
import math
import os
from pathlib import Path
import tempfile
import tomllib
from zipfile import ZipFile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Inches, Pt


def _number(value, name, *, positive=False):
    if type(value) not in (int, float) or not math.isfinite(value) or value < 0 or (positive and value == 0):
        raise ValueError(f'{name} must be finite and {"positive" if positive else "nonnegative"}')
    return value


def _text(value):
    if not isinstance(value, str):
        raise ValueError('text values must be strings')
    return value


def _color(value):
    if not isinstance(value, str) or len(value) != 7 or not value.startswith('#'):
        raise ValueError('colors must be #RRGGBB')
    return RGBColor.from_string(value[1:])


def _style(frame, text, points, *, bold=False, color='#000000'):
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = Inches(.08)
    frame.margin_top = frame.margin_bottom = Inches(.04)
    for index, line in enumerate(text.split('\n')):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.name = 'Meiryo'
        run.font.size = Pt(points)
        run.font.bold = bold
        run.font.color.rgb = _color(color)


def _shape(slide, kind, rect, name, *, fill='#FFFFFF', line='#000000'):
    shape = slide.shapes.add_shape(kind, *(Inches(v) for v in rect))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _color(fill)
    shape.line.color.rgb = _color(line)
    shape.line.width = Pt(1)
    return shape


def _textbox(slide, rect, text, name, points, *, bold=False):
    shape = slide.shapes.add_textbox(*(Inches(v) for v in rect))
    shape.name = name
    _style(shape.text_frame, text, points, bold=bold)
    return shape


def enrich(source, patterns, *, destination=None, dry_run=False):
    """Append only; copy every non-target ZIP member without modification.

    Coordinates are inches. Callers choose a free region; overlap with existing
    objects is not resolved. Dry-run builds the same native objects in memory.
    """
    if not isinstance(patterns, list) or not patterns:
        raise ValueError('at least one pattern is required')
    if not dry_run and (destination is None or Path(source).resolve() == Path(destination).resolve() or os.path.lexists(destination)):
        raise ValueError('destination must be a new file distinct from source')
    prs = Presentation(str(source))
    changed, report = {}, []
    for index, spec in enumerate(patterns):
        if not isinstance(spec, dict):
            raise ValueError('pattern must be a mapping')
        mode, number = spec.get('pattern'), spec.get('slide')
        if mode not in ('cards', 'regions', 'agenda', 'flow', 'table'):
            raise ValueError('unknown pattern')
        if type(number) is not int or not 1 <= number <= len(prs.slides):
            raise ValueError('slide must be a valid 1-based slide number')
        rect = [_number(spec.get(k), k, positive=k in ('width','height')) for k in ('x','y','width','height')]
        x,y,w,h = rect
        if Inches(x+w) > prs.slide_width or Inches(y+h) > prs.slide_height:
            raise ValueError('pattern extends beyond slide bounds')
        gap = _number(spec.get('gap', .2), 'gap')
        points = _number(spec.get('font_size', 14), 'font_size', positive=True)
        if not 6 <= points <= 72:
            raise ValueError('font_size must be between 6 and 72')
        accent = spec.get('accent', '#000000')
        _color(accent)
        slide = prs.slides[number-1]
        before = len(slide.shapes)
        prefix = f'enrich-{index+1}-{mode}'
        if mode == 'table':
            if 'csv' in spec and 'rows' in spec:
                raise ValueError('use either csv or rows, not both')
            if 'csv' in spec:
                with Path(_text(spec['csv'])).open(encoding='utf-8-sig', newline='') as stream:
                    rows = list(csv.reader(stream))
            else:
                rows = spec.get('rows')
            header_fill = spec.get('header_fill', '#222222')
            header_text = spec.get('header_text', '#FFFFFF')
            row_fills = spec.get('row_fills', ['#FFFFFF', '#F2F2F2'])
            border = spec.get('border_color', '#B3B3B3')
            if not isinstance(row_fills, list) or not row_fills:
                raise ValueError('row_fills must be a nonempty list of colors')
            for value in [header_fill, header_text, border, *row_fills]:
                _color(value)
            if (not isinstance(rows,list) or not 2 <= len(rows) <= 30 or not isinstance(rows[0],list)
                    or not 1 <= len(rows[0]) <= 12 or any(not isinstance(r,list) or len(r)!=len(rows[0]) for r in rows)):
                raise ValueError('table requires 2..30 rectangular rows with 1..12 columns')
            for row in rows:
                for value in row:
                    _text(value)
            shape = slide.shapes.add_table(len(rows),len(rows[0]),*(Inches(v) for v in rect))
            shape.name = prefix
            table = shape.table
            for ri,row in enumerate(rows):
                for ci,value in enumerate(row):
                    cell = table.cell(ri,ci)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _color(header_fill if ri==0 else row_fills[(ri-1)%len(row_fills)])
                    _style(cell.text_frame,value,points,bold=ri==0,color=header_text if ri==0 else '#000000')
                    props = cell._tc.get_or_add_tcPr()
                    for side in ('L', 'R', 'T', 'B'):
                        tag = 'a:ln'+side
                        line = OxmlElement(tag)
                        line.set('w', '6350') # 0.5 points
                        fill = OxmlElement('a:solidFill')
                        rgb = OxmlElement('a:srgbClr'); rgb.set('val', border[1:].upper())
                        fill.append(rgb); line.append(fill)
                        props.append(line)
        else:
            items = spec.get('items')
            if not isinstance(items,list) or not 1 <= len(items) <= 24:
                raise ValueError('items must contain 1..24 mappings')
            for item in items:
                if not isinstance(item,dict):
                    raise ValueError('item must be a mapping')
                _text(item.get('title',''))
                _text(item.get('body',''))
                _text(item.get('label',''))
            n = len(items)
            columns = spec.get('columns', n)
            if mode in ('regions','flow'):
                columns = n
            elif mode == 'agenda':
                columns = 1
            if type(columns) is not int or not 1 <= columns <= n:
                raise ValueError('columns must be between 1 and item count')
            row_count = (n+columns-1)//columns
            cw,ch = (w-gap*(columns-1))/columns,(h-gap*(row_count-1))/row_count
            if cw < .7 or ch < .7:
                raise ValueError('cells must be at least 0.7 inches wide and high')
            if mode == 'flow' and n>1 and gap < .15:
                raise ValueError('flow gap must be at least 0.15 inches')
            for i,item in enumerate(items):
                cx,cy = x+(i%columns)*(cw+gap), y+(i//columns)*(ch+gap)
                name = f'{prefix}-{i+1}'
                kind = MSO_SHAPE.ROUNDED_RECTANGLE if mode in ('regions','agenda') else MSO_SHAPE.RECTANGLE
                _shape(slide,kind,(cx,cy,cw,ch),name+'-box',line=accent)
                if mode == 'agenda':
                    _textbox(slide,(cx+.05,cy+.03,cw-.1,ch-.06),item.get('title','')+'\n'+item.get('body','')+('  '+item['label'] if item.get('label') else ''),name+'-text',points,bold=False)
                else:
                    header = min(.45,ch*.3)
                    if mode == 'regions':
                        badge = _shape(slide,MSO_SHAPE.ROUNDED_RECTANGLE,(cx+.05,cy+.03,cw-.1,header),name+'-label',fill=accent,line=accent)
                        _style(badge.text_frame,item.get('title',''),points,bold=True,color='#FFFFFF')
                    else:
                        _shape(slide,MSO_SHAPE.RECTANGLE,(cx,cy,cw,.035),name+'-bar',fill=accent,line=accent)
                        _textbox(slide,(cx+.04,cy+.04,cw-.08,header),item.get('title',''),name+'-title',points,bold=True)
                    _textbox(slide,(cx+.04,cy+header+.07,cw-.08,ch-header-.1),item.get('body',''),name+'-body',points)
                if mode == 'flow' and i<n-1:
                    arrow = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                        Inches(cx+cw+gap*.15), Inches(cy+ch/2),
                        Inches(cx+cw+gap*.85), Inches(cy+ch/2))
                    arrow.name = name+'-arrow'
                    arrow.line.color.rgb = _color(accent)
                    arrow.line.width = Pt(1.5)
                    end = OxmlElement('a:tailEnd')
                    end.set('type', 'triangle')
                    end.set('w', 'med'); end.set('len', 'med')
                    arrow.line._get_or_add_ln().append(end)
        member = str(slide.part.partname).lstrip('/')
        changed[member] = slide.part.blob
        report.append(dict(pattern=mode,slide=number,added_shapes=len(slide.shapes)-before,
                           shape_ids=[s.shape_id for s in list(slide.shapes)[before:]]))
    with ZipFile(source) as original:
        if len(original.namelist()) != len(set(original.namelist())):
            raise ValueError('duplicate ZIP members unsupported')
        if not dry_run:
            dst = Path(destination)
            with tempfile.TemporaryDirectory(prefix='.pptx-enrich-', dir=dst.parent) as staging:
                staged = Path(staging)/'complete.pptx'
                with ZipFile(staged,'w') as output:
                    output.comment = original.comment
                    for entry in original.infolist():
                        output.writestr(entry, changed.get(entry.filename, original.read(entry.filename)))
                Presentation(staged)
                os.link(staged,dst)
    return dict(dry_run=dry_run,patterns=report,added_shapes=sum(r['added_shapes'] for r in report))


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source',type=Path)
    parser.add_argument('--patterns',type=Path,required=True)
    parser.add_argument('--output',type=Path)
    parser.add_argument('--dry-run',action='store_true')
    args = parser.parse_args(argv)
    with args.patterns.open('rb') as stream:
        patterns = tomllib.load(stream).get('patterns')
    if isinstance(patterns, list):
        for spec in patterns:
            if isinstance(spec, dict) and isinstance(spec.get('csv'), str):
                spec['csv'] = str(args.patterns.parent / spec['csv'])
    print(json.dumps(enrich(args.source,patterns,destination=args.output,dry_run=args.dry_run),ensure_ascii=False,indent=2))


if __name__ == '__main__':
    main()
