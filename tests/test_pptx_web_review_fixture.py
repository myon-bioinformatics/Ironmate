"""Web review fixture must demonstrate a real native Open XML edit."""
from pathlib import Path
from pptx import Presentation
from pptx_web_review_fixture import generate


def test_before_after_review_fixture(tmp_path):
    report = generate(tmp_path)
    assert report["changed_members"] == ["ppt/slides/slide1.xml"]
    assert set(report["files"]) == {"native-before.pptx", "native-after.pptx"}
    assert (tmp_path / "review-manifest.json").is_file()
    before = Presentation(tmp_path / "native-before.pptx").slides[0].shapes[0]
    after = Presentation(tmp_path / "native-after.pptx").slides[0].shapes[0]
    assert before.left != after.left
    assert before.fill.fore_color.rgb != after.fill.fore_color.rgb
    assert before.text == after.text
