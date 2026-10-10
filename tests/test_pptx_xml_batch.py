"""Multi-slide Open XML mapping, selective changes, and atomic rejection."""
from zipfile import ZipFile

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from pptx_xml_batch import edit_shape_fills


def make_deck(path, *, reordered=False):
    presentation = Presentation()
    for i in range(3):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1),
                                       Inches(3), Inches(1))
        shape.name = "Shared"
        shape.text = f"Slide {i + 1}"
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(10, 20, 30)
    if reordered:
        ids = presentation.slides._sldIdLst
        first = ids[0]
        ids.remove(first)
        ids.append(first)
    presentation.save(path)


def rgb(deck, index):
    return Presentation(deck).slides[index].shapes[0].fill.fore_color.rgb


def test_selective_and_all_slide_recolor(tmp_path):
    source, selective, all_slides = (tmp_path / n for n in ("original.pptx", "selective.pptx", "all.pptx"))
    make_deck(source)
    report = edit_shape_fills(source, selective, [
        {"slide": 2, "shape_name": "Shared", "fill_rgb": "#AABBCC"}])
    assert report["changed_slides"] == [2]
    assert [rgb(selective, i) for i in range(3)] == [
        RGBColor(10, 20, 30), RGBColor(170, 187, 204), RGBColor(10, 20, 30)]
    rules = [{"slide": i, "shape_name": "Shared", "fill_rgb": "#FFFFFF"} for i in range(1, 4)]
    result = edit_shape_fills(selective, all_slides, rules)
    assert result["changed_slides"] == [1, 2, 3]
    assert all(rgb(all_slides, i) == RGBColor(255, 255, 255) for i in range(3))
    with ZipFile(source) as old, ZipFile(selective) as new:
        assert old.namelist() == new.namelist()
        assert [name for name in old.namelist() if old.read(name) != new.read(name)] == [
            "ppt/slides/slide2.xml"]


def test_reordered_slide_uses_display_order_not_xml_suffix(tmp_path):
    source, output = tmp_path / "reordered.pptx", tmp_path / "edited.pptx"
    make_deck(source, reordered=True)
    assert [s.shapes[0].text for s in Presentation(source).slides] == [
        "Slide 2", "Slide 3", "Slide 1"]
    report = edit_shape_fills(source, output, [
        {"slide": 1, "shape_name": "Shared", "fill_rgb": "#ABCDEF"}])
    # Raw presentation.xml.rels maps the first displayed slide to slide2.xml.
    # python-pptx partname can be re-normalized and is not a stored ZIP locator.
    assert report["changed_members"] == ["ppt/slides/slide2.xml"]
    assert rgb(output, 0) == RGBColor(171, 205, 239)
    assert rgb(output, 1) == RGBColor(10, 20, 30)


@pytest.mark.parametrize("rules", [
    [{"slide": 4, "shape_name": "Shared", "fill_rgb": "#FFFFFF"}],
    [{"slide": 1, "shape_name": "Missing", "fill_rgb": "#FFFFFF"}],
    [{"slide": 1, "shape_name": "Shared", "fill_rgb": "white"}],
    [{"slide": 1, "shape_name": "Shared", "fill_rgb": "#FFFFFF"},
     {"slide": 1, "shape_name": "Shared", "fill_rgb": "#000000"}],
    [{"slide": 1, "shape_name": "Shared", "fill_rgb": "#FFFFFF"},
     {"slide": 3, "shape_name": "Missing", "fill_rgb": "#000000"}],
])
def test_invalid_batch_writes_nothing(tmp_path, rules):
    source, output = tmp_path / "source.pptx", tmp_path / "edited.pptx"
    make_deck(source)
    with pytest.raises(ValueError):
        edit_shape_fills(source, output, rules)
    assert not output.exists()
    assert all(rgb(source, i) == RGBColor(10, 20, 30) for i in range(3))
