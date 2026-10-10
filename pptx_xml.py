"""Open XML edits: slide metadata, native text/shapes, and multi-slide batches.

All functions are experimental and write a distinct output PPTX.
"""

# --- pptx_xml_probe.py ---
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from pptx import Presentation

P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def set_slide_name(source, destination, *, slide=1, name="Diagram"):
    """Edit p:cSld/@name in a new PPTX; never rewrite the source."""
    source, destination = Path(source), Path(destination)
    if not isinstance(name, str) or not name.strip():
        raise ValueError("name must be nonempty")
    if type(slide) is not int or slide < 1:
        raise ValueError("slide must be a positive integer")
    if source.resolve() == destination.resolve() or destination.exists():
        raise ValueError("destination must be a new file")
    member = f"ppt/slides/slide{slide}.xml"
    with ZipFile(source) as original:
        names = original.namelist()
        if member not in names:
            raise ValueError("slide XML member not found")
        if len(names) != len(set(names)):
            raise ValueError("duplicate ZIP members are unsupported")
        root = ET.fromstring(original.read(member))
        csld = root.find(P + "cSld")
        if csld is None:
            raise ValueError("p:cSld missing")
        csld.set("name", name)
        replacement = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        try:
            with ZipFile(destination, "x") as changed:
                for entry in original.infolist():
                    changed.writestr(entry, replacement if entry.filename == member
                                     else original.read(entry.filename))
        except BaseException:
            destination.unlink(missing_ok=True)
            raise
    Presentation(destination)
    return {"source": str(source), "output": str(destination), "member": member, "name": name}



A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def set_run_style(source, destination, *, slide=1, rgb=None, points=None, bold=None):
    """Experimental native a:rPr edit; reject inherited colors and image text."""
    if rgb is None and points is None and bold is None:
        raise ValueError("at least one style attribute required")
    if rgb is not None:
        if not isinstance(rgb, str) or len(rgb) != 7 or rgb[0] != "#" or not all(
                char in "0123456789abcdefABCDEF" for char in rgb[1:]):
            raise ValueError("rgb must be #RRGGBB")
    if points is not None and (type(points) not in (int, float) or not 1 <= points <= 400):
        raise ValueError("points must be between 1 and 400")
    if bold is not None and type(bold) is not bool:
        raise ValueError("bold must be boolean")
    if type(slide) is not int or slide < 1:
        raise ValueError("slide must be positive")
    source, destination = Path(source), Path(destination)
    if source.resolve() == destination.resolve() or destination.exists():
        raise ValueError("destination must be a new file")
    member = f"ppt/slides/slide{slide}.xml"
    with ZipFile(source) as original:
        names = original.namelist()
        if member not in names or len(names) != len(set(names)):
            raise ValueError("missing or duplicate slide XML member")
        root = ET.fromstring(original.read(member))
        runs = root.findall(".//" + A + "r")
        if not runs:
            raise ValueError("no editable native text runs")
        # Prevalidate every run before modifying any XML or writing output.
        if rgb is not None:
            for run in runs:
                props = run.find(A + "rPr")
                fill = props.find(A + "solidFill") if props is not None else None
                if fill is None or len(fill) != 1 or fill[0].tag != A + "srgbClr":
                    raise ValueError("explicit RGB run color required")
        for run in runs:
            props = run.find(A + "rPr")
            if props is None:
                props = ET.Element(A + "rPr")
                run.insert(0, props)
            if points is not None:
                props.set("sz", str(round(points * 100)))
            if bold is not None:
                props.set("b", "1" if bold else "0")
            if rgb is not None:
                props.find(A + "solidFill")[0].set("val", rgb[1:].upper())
        replacement = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        try:
            with ZipFile(destination, "x") as changed:
                for entry in original.infolist():
                    changed.writestr(entry, replacement if entry.filename == member
                                     else original.read(entry.filename))
            Presentation(destination)
        except BaseException:
            destination.unlink(missing_ok=True)
            raise
    return {"output": str(destination), "member": member, "runs": len(runs)}


# --- pptx_xml_shapes.py ---
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from pptx import Presentation
from pptx.util import Emu

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def edit_shape(source, destination, *, slide=1, shape_name, fill_rgb=None, line_rgb=None,
               left=None, top=None, width=None, height=None):
    if not isinstance(shape_name, str) or not shape_name.strip():
        raise ValueError("shape_name is required")
    if type(slide) is not int or slide < 1:
        raise ValueError("invalid slide")
    if all(value is None for value in (fill_rgb, line_rgb, left, top, width, height)):
        raise ValueError("no changes requested")
    for color in (fill_rgb, line_rgb):
        if color is not None and (not isinstance(color, str) or len(color) != 7
                                  or color[0] != "#" or any(c not in "0123456789abcdefABCDEF" for c in color[1:])):
            raise ValueError("colors must be #RRGGBB")
    for key, value in (("left", left), ("top", top), ("width", width), ("height", height)):
        if value is not None and (not isinstance(value, int) or isinstance(value, bool) or (key in ("width", "height") and value <= 0)):
            raise ValueError("geometry must be integer EMUs with positive size")
    source, destination = Path(source), Path(destination)
    if source.resolve() == destination.resolve() or destination.exists():
        raise ValueError("destination must be new and distinct")
    member = f"ppt/slides/slide{slide}.xml"
    with ZipFile(source) as package:
        entries = package.infolist()
        names = [entry.filename for entry in entries]
        if len(names) != len(set(names)) or member not in names:
            raise ValueError("invalid or missing slide XML")
        root = ET.fromstring(package.read(member))
        matches = [node for node in root.findall(".//" + P + "sp")
                   if (node.find(P + "nvSpPr/" + P + "cNvPr") is not None
                       and node.find(P + "nvSpPr/" + P + "cNvPr").get("name") == shape_name)]
        if len(matches) != 1:
            raise ValueError("expected exactly one named native shape")
        shape = matches[0]
        props = shape.find(P + "spPr")
        if props is None:
            raise ValueError("shape properties missing")
        if fill_rgb is not None:
            fill = props.find(A + "solidFill")
            if fill is None or len(fill) != 1 or fill[0].tag != A + "srgbClr":
                raise ValueError("explicit RGB shape fill required")
            fill[0].set("val", fill_rgb[1:].upper())
        if line_rgb is not None:
            line = props.find(A + "ln")
            fill = line.find(A + "solidFill") if line is not None else None
            if fill is None or len(fill) != 1 or fill[0].tag != A + "srgbClr":
                raise ValueError("explicit RGB shape line required")
            fill[0].set("val", line_rgb[1:].upper())
        if any(value is not None for value in (left, top, width, height)):
            xfrm = props.find(A + "xfrm")
            if xfrm is None or xfrm.find(A + "off") is None or xfrm.find(A + "ext") is None:
                raise ValueError("explicit shape geometry required")
            off, ext = xfrm.find(A + "off"), xfrm.find(A + "ext")
            for node, key, value in ((off, "x", left), (off, "y", top),
                                     (ext, "cx", width), (ext, "cy", height)):
                if value is not None:
                    node.set(key, str(value))
        replacement = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        try:
            with ZipFile(destination, "x") as changed:
                for entry in entries:
                    changed.writestr(entry, replacement if entry.filename == member
                                     else package.read(entry.filename))
            Presentation(destination)
        except BaseException:
            destination.unlink(missing_ok=True)
            raise
    return {"output": str(destination), "shape_name": shape_name, "slide": slide}



# --- pptx_xml_batch.py ---
from pathlib import Path, PurePosixPath
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from pptx import Presentation

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
REL = "{http://schemas.openxmlformats.org/package/2006/relationships}"


def _ordered_slide_members(package):
    """Map display-order r:ids to the *stored* ZIP members, without OPC reserialization."""
    presentation = ET.fromstring(package.read("ppt/presentation.xml"))
    rels = ET.fromstring(package.read("ppt/_rels/presentation.xml.rels"))
    mapping = {}
    for rel in rels.findall(REL + "Relationship"):
        if rel.get("Type", "").endswith("/slide"):
            if rel.get("TargetMode", "Internal") != "Internal":
                raise ValueError("external slide relationships are unsupported")
            rid, target = rel.get("Id"), rel.get("Target", "")
            if not rid or not target or ".." in PurePosixPath(target).parts or target.startswith("/"):
                raise ValueError("invalid slide relationship target")
            if rid in mapping:
                raise ValueError("duplicate slide relationship")
            mapping[rid] = "ppt/" + str(PurePosixPath(target))
    ids = presentation.find(P + "sldIdLst")
    if ids is None:
        raise ValueError("slide order list missing")
    ordered = []
    for slide in ids.findall(P + "sldId"):
        rid = slide.get(R)
        if rid not in mapping:
            raise ValueError("unresolved slide relationship")
        ordered.append(mapping[rid])
    if len(ordered) != len(set(ordered)):
        raise ValueError("duplicate slide target")
    return ordered


def edit_shape_fills(source, destination, rules):
    """Rules: list of {slide: 1-based display index, shape_name: str, fill_rgb: #RRGGBB}."""
    source, destination = Path(source), Path(destination)
    if source.resolve() == destination.resolve() or destination.exists():
        raise ValueError("destination must be a new, distinct file")
    if not isinstance(rules, list) or not rules:
        raise ValueError("nonempty rules list required")
    presentation = Presentation(source)
    parsed = {}
    targets = set()
    with ZipFile(source) as package:
        entries = package.infolist()
        names = [entry.filename for entry in entries]
        parts = _ordered_slide_members(package)
        if len(parts) != len(presentation.slides):
            raise ValueError("slide count does not match presentation order")
        if len(names) != len(set(names)) or any(part not in names for part in parts):
            raise ValueError("invalid slide package members")
        for rule in rules:
            if not isinstance(rule, dict) or set(rule) != {"slide", "shape_name", "fill_rgb"}:
                raise ValueError("rule requires slide, shape_name, fill_rgb")
            index, shape_name, rgb = rule["slide"], rule["shape_name"], rule["fill_rgb"]
            if type(index) is not int or not 1 <= index <= len(parts):
                raise ValueError("slide index out of range")
            if not isinstance(shape_name, str) or not shape_name.strip():
                raise ValueError("nonempty shape_name required")
            if not isinstance(rgb, str) or len(rgb) != 7 or not rgb.startswith("#") or any(
                    ch not in "0123456789abcdefABCDEF" for ch in rgb[1:]):
                raise ValueError("fill_rgb must be #RRGGBB")
            key = (index, shape_name)
            if key in targets:
                raise ValueError("duplicate edit target")
            targets.add(key)
            member = parts[index - 1]
            if member not in parsed:
                parsed[member] = ET.fromstring(package.read(member))
            root = parsed[member]
            shapes = [
                shape for shape in root.findall(".//" + P + "sp")
                if shape.find(P + "nvSpPr/" + P + "cNvPr") is not None
                and shape.find(P + "nvSpPr/" + P + "cNvPr").get("name") == shape_name
            ]
            if len(shapes) != 1:
                raise ValueError("expected exactly one named native shape per slide")
            props = shapes[0].find(P + "spPr")
            fill = props.find(A + "solidFill") if props is not None else None
            if fill is None or len(fill) != 1 or fill[0].tag != A + "srgbClr":
                raise ValueError("explicit RGB shape fill required")
            fill[0].set("val", rgb[1:].upper())
        replacements = {
            member: ET.tostring(root, encoding="utf-8", xml_declaration=True)
            for member, root in parsed.items()
        }
        try:
            with ZipFile(destination, "x") as output:
                for entry in entries:
                    output.writestr(entry, replacements.get(entry.filename, package.read(entry.filename)))
            check = Presentation(destination)
            if len(check.slides) != len(parts):
                raise ValueError("slide count changed")
        except BaseException:
            destination.unlink(missing_ok=True)
            raise
    return {"changed_members": sorted(replacements), "changed_slides": sorted({key[0] for key in targets})}
