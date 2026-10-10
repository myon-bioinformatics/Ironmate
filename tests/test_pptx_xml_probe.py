"""Open XML metadata mutation must preserve unrelated ZIP parts and reopen."""
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from pptx import Presentation
from pptx_xml import set_slide_name
import pytest


def test_slide_name_xml_edit_roundtrip(tmp_path):
    source = tmp_path / "original.pptx"
    destination = tmp_path / "edited.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(source)
    result = set_slide_name(source, destination, name="UML diagram")
    assert result["member"] == "ppt/slides/slide1.xml"
    with ZipFile(source) as before, ZipFile(destination) as after:
        assert before.namelist() == after.namelist()
        for name in before.namelist():
            if name != "ppt/slides/slide1.xml":
                assert before.read(name) == after.read(name)
        root = ET.fromstring(after.read("ppt/slides/slide1.xml"))
        ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
        assert root.find("p:cSld", ns).get("name") == "UML diagram"
    assert len(Presentation(destination).slides) == 1


def test_xml_edit_rejects_overwrite(tmp_path):
    source = tmp_path / "original.pptx"
    prs = Presentation()
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(source)
    with pytest.raises(ValueError, match="new file"):
        set_slide_name(source, source)




def test_xml_native_run_style_roundtrip(tmp_path):
    from pptx.dml.color import RGBColor
    from pptx.util import Inches
    from pptx_xml import set_run_style
    src, dst = tmp_path / "source.pptx", tmp_path / "style.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Editable"
    run.font.color.rgb = RGBColor(100, 100, 100)
    prs.save(src)
    report = set_run_style(src, dst, rgb="#112233", points=24, bold=True)
    assert report["runs"] == 1
    actual = Presentation(dst).slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert actual.font.color.rgb == RGBColor(17, 34, 51)
    assert actual.font.size.pt == 24
    assert actual.font.bold is True
    assert actual.text == "Editable"
    with ZipFile(src) as before, ZipFile(dst) as after:
        assert before.namelist() == after.namelist()
        assert all(before.read(name) == after.read(name) for name in before.namelist()
                   if name != "ppt/slides/slide1.xml")


def test_xml_style_rejects_inherited_rgb_without_output(tmp_path):
    from pptx.util import Inches
    from pptx_xml import set_run_style
    src, dst = tmp_path / "source.pptx", tmp_path / "style.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Inherited"
    prs.save(src)
    with pytest.raises(ValueError, match="explicit RGB"):
        set_run_style(src, dst, rgb="#000000")
    assert not dst.exists()


@pytest.mark.parametrize("kwargs", [
    {"rgb": "#xyzxyz"}, {"points": 0}, {"bold": "true"}, {}
])
def test_xml_style_invalid_rules_fail_closed(tmp_path, kwargs):
    from pptx_xml import set_run_style
    with pytest.raises(ValueError):
        set_run_style(tmp_path / "missing.pptx", tmp_path / "out.pptx", **kwargs)
