"""Test-only prototype: background line and separate angular arrowheads."""
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def add_diagram(slide, *, chevron=False):
    # Geometry in inches: full line extends behind both boxes.
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(0.5), Inches(1.4), Inches(9.5), Inches(1.4))
    line.line.color.rgb = RGBColor(0, 0, 0)
    line.line.width = Pt(2)
    boxes = []
    for x, text in ((0.5, "A"), (7.5, "B")):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(0.9), Inches(2), Inches(1))
        shape.text = text
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        boxes.append(shape)
    head = slide.shapes.add_shape(
        MSO_SHAPE.CHEVRON if chevron else MSO_SHAPE.RIGHT_TRIANGLE,
        Inches(4.75), Inches(1.22), Inches(0.35), Inches(0.36))
    head.fill.solid()
    head.fill.fore_color.rgb = RGBColor(0, 0, 0)
    head.line.fill.background()
    return line, boxes, head


def test_diagram_arrowhead_variants(tmp_path):
    for chevron in (False, True):
        p = Presentation()
        slide = p.slides.add_slide(p.slide_layouts[6])
        line, boxes, head = add_diagram(slide, chevron=chevron)
        assert slide.shapes[0] == line
        assert slide.shapes[-1] == head
        assert line.left == boxes[0].left
        assert line.left + line.width == boxes[1].left + boxes[1].width
        assert boxes[0].left + boxes[0].width < head.left < boxes[1].left
        path = tmp_path / ("chevron.pptx" if chevron else "triangle.pptx")
        p.save(path)
        assert len(Presentation(path).slides[0].shapes) == 4


def test_multiple_horizontal_boxes_have_independent_midpoint_heads(tmp_path):
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[6])
    boxes = []
    for x in (1, 4, 7):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(1), Inches(1), Inches(1))
        boxes.append(shape)
    for first, second in zip(boxes, boxes[1:]):
        center_y = first.top + first.height // 2
        line = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT, first.left, center_y,
            second.left + second.width, center_y)
        midpoint = (first.left + first.width + second.left) // 2
        head = slide.shapes.add_shape(
            MSO_SHAPE.RIGHT_TRIANGLE, midpoint - Inches(0.15),
            center_y - Inches(0.15), Inches(0.3), Inches(0.3))
        assert head.left + head.width // 2 == midpoint
        assert line.left == first.left
        assert line.left + line.width == second.left + second.width
    path = tmp_path / "horizontal.pptx"
    p.save(path)
    assert len(Presentation(path).slides[0].shapes) == 7


def test_vertical_boxes_use_full_height_line_and_midpoint_head(tmp_path):
    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[6])
    upper = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(3), Inches(1), Inches(2), Inches(1))
    lower = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(3), Inches(4), Inches(2), Inches(1))
    center_x = upper.left + upper.width // 2
    line = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, center_x, upper.top,
        center_x, lower.top + lower.height)
    midpoint = (upper.top + upper.height + lower.top) // 2
    head = slide.shapes.add_shape(
        MSO_SHAPE.DOWN_ARROW, center_x - Inches(0.15),
        midpoint - Inches(0.15), Inches(0.3), Inches(0.3))
    assert head.top + head.height // 2 == midpoint
    assert line.top == upper.top
    assert line.top + line.height == lower.top + lower.height
    path = tmp_path / "vertical.pptx"
    p.save(path)
    assert len(Presentation(path).slides[0].shapes) == 4


def edge_center(shape, edge):
    if edge == "left":
        return shape.left, shape.top + shape.height // 2
    if edge == "right":
        return shape.left + shape.width, shape.top + shape.height // 2
    if edge == "top":
        return shape.left + shape.width // 2, shape.top
    if edge == "bottom":
        return shape.left + shape.width // 2, shape.top + shape.height
    raise ValueError("edge must be left/right/top/bottom")


def test_diagonal_edge_selection_and_rotation(tmp_path):
    import math

    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[6])
    a = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(2), Inches(1))
    c = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(6), Inches(4), Inches(2), Inches(1))
    for start_edge in ("left", "top", "bottom", "right"):
        start = edge_center(c, start_edge)
        end = edge_center(a, "bottom")
        dx, dy = end[0] - start[0], end[1] - start[1]
        angle = math.degrees(math.atan2(dy, dx))
        assert math.isfinite(angle)
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, *start, *end)
        mid_x, mid_y = (start[0] + end[0]) // 2, (start[1] + end[1]) // 2
        head = slide.shapes.add_shape(
            MSO_SHAPE.RIGHT_TRIANGLE, mid_x - Inches(0.15),
            mid_y - Inches(0.15), Inches(0.3), Inches(0.3))
        head.rotation = angle
        assert abs((head.rotation - angle + 180) % 360 - 180) < 0.01
        assert line is not None
    with pytest.raises(ValueError, match="edge"):
        edge_center(c, "diagonal")
    path = tmp_path / "diagonal.pptx"
    p.save(path)
    assert len(Presentation(path).slides[0].shapes) == 10
