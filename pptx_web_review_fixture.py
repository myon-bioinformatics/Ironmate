"""Produce before/after native PPTX fixtures for PowerPoint Web visual review.

This script does not claim that PowerPoint Web has been opened or verified.
"""
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches

from pptx_xml import set_run_style
from pptx_xml import edit_shape


def generate(directory):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    original = directory / "native-before.pptx"
    text_edited = directory / "native-text-edited.pptx"
    final = directory / "native-after.pptx"
    if any(path.exists() for path in (original, text_edited, final)):
        raise ValueError("output fixtures already exist")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(2), Inches(3), Inches(1.3))
    box.name = "Target"
    box.fill.solid()
    box.fill.fore_color.rgb = RGBColor(230, 235, 245)
    box.line.color.rgb = RGBColor(50, 70, 90)
    run = box.text_frame.paragraphs[0].add_run()
    run.text = "Open XML edit"
    run.font.color.rgb = RGBColor(20, 30, 40)
    run.font.size = Inches(0.22)
    presentation.save(original)

    set_run_style(original, text_edited, rgb="#FFFFFF", points=22, bold=True)
    edit_shape(text_edited, final, shape_name="Target", fill_rgb="#14213D",
               line_rgb="#FFFFFF", left=Inches(2), top=Inches(2),
               width=Inches(4), height=Inches(1.5))

    before = Presentation(original).slides[0].shapes[0]
    after = Presentation(final).slides[0].shapes[0]
    assert before.text == after.text == "Open XML edit"
    assert after.fill.fore_color.rgb == RGBColor(20, 33, 61)
    assert after.line.color.rgb == RGBColor(255, 255, 255)
    assert after.text_frame.paragraphs[0].runs[0].font.color.rgb == RGBColor(255, 255, 255)
    assert after.text_frame.paragraphs[0].runs[0].font.bold is True
    assert after.text_frame.paragraphs[0].runs[0].font.size.pt == 22

    with ZipFile(original) as old, ZipFile(final) as new:
        names = old.namelist()
        assert names == new.namelist()
        changed = [name for name in names if old.read(name) != new.read(name)]
        assert changed == ["ppt/slides/slide1.xml"]
    report = {
        "verification": "PPTX structure only; PowerPoint Web rendering not yet observed",
        "changed_members": changed,
        "files": {
            path.name: {"sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                        "bytes": path.stat().st_size}
            for path in (original, final)
        },
        "expected": {"before_fill": "#E6EBF5", "after_fill": "#14213D",
                     "after_text": "#FFFFFF", "after_line": "#FFFFFF",
                     "after_bold": True, "after_font_points": 22},
    }
    (directory / "review-manifest.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return report


if __name__ == "__main__":
    print(json.dumps(generate("build/test-results/pptx-web-review"), indent=2))
