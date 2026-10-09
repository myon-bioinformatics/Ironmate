import subprocess
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor

from markdown_pptx import generate, parse_blocks
from pptx_replace import transform


def test_markdown_pptx_xml_edit_roundtrip(tmp_path):
    source = Path(__file__).parent / "fixtures" / "mini_readme.md"
    original, edited = tmp_path / "original.pptx", tmp_path / "edited.pptx"
    manifest = generate(source, original)
    assert len(manifest["blocks"]) == 5
    prs = Presentation(original)
    shapes = list(prs.slides[0].shapes)
    assert [s.name for s in shapes] == [b["shape_name"] for b in manifest["blocks"]]
    assert [s.text for s in shapes] == [b["text"] for b in manifest["blocks"]]
    assert all(s.text_frame.paragraphs[0].runs[0].font.color.rgb == RGBColor(0, 0, 0) for s in shapes)

    result = transform(original, [{"operation": "text_color", "color": "#2244AA"}], destination=edited)
    assert result["change_count"] == 5
    changed = Presentation(edited).slides[0].shapes
    assert [s.text for s in changed] == [b["text"] for b in manifest["blocks"]]
    assert all(s.text_frame.paragraphs[0].runs[0].font.color.rgb == RGBColor(34, 68, 170) for s in changed)

    with zipfile.ZipFile(original) as before, zipfile.ZipFile(edited) as after:
        names = set(before.namelist())
        assert names == set(after.namelist())
        for name in names - {"ppt/slides/slide1.xml"}:
            assert before.read(name) == after.read(name), name
        ns = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
        xml = ET.fromstring(after.read("ppt/slides/slide1.xml"))
        assert [t.text for t in xml.findall(".//a:t", ns)] == [b["text"] for b in manifest["blocks"]]


def test_markdown_cli_one_liner(tmp_path):
    source = Path(__file__).parent / "fixtures" / "mini_readme.md"
    output = tmp_path / "fixture.pptx"
    script = Path(__file__).resolve().parents[1] / "markdown_pptx.py"
    p = subprocess.run([sys.executable, str(script), str(source), "-o", str(output)], capture_output=True, text=True)
    assert p.returncode == 0, p.stderr
    assert output.is_file()
    assert "markdown:line-1:title" in p.stdout


@pytest.mark.parametrize("text", ["", "### Unsupported", "- list", "# ", "\n".join(["text"] * 9)])
def test_markdown_explicit_failures(text):
    with pytest.raises(ValueError):
        parse_blocks(text)


def test_markdown_output_collision(tmp_path):
    source = Path(__file__).parent / "fixtures" / "mini_readme.md"
    output = tmp_path / "existing.pptx"
    output.write_bytes(b"unchanged")
    with pytest.raises(ValueError, match="new file"):
        generate(source, output)
    assert output.read_bytes() == b"unchanged"
