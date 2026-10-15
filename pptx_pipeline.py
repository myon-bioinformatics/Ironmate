"""One-command Markdown -> Marp PPTX -> Open XML inspection.

No promise of Mermaid rendering unless a configured Marp renderer supports it.
"""
import json
import csv
import math
import tomllib
from zipfile import ZipFile
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Inches, Pt
from vendor.markdown import markdown_table_to_rows
import os
from pathlib import Path
import sys
import tempfile

from vendor.cli_args import Argument, make_parser
import shutil
import subprocess
from pptx import Presentation
from mermaid_svg import render_mermaid
from pptx_inspect import inspect



def _number(value, name, *, positive=False):
    if type(value) not in (int, float) or not math.isfinite(value) or value < 0 or (positive and value == 0):
        raise ValueError(f'{name} must be finite and {"positive" if positive else "nonnegative"}')
    return value


def _text(value):
    if not isinstance(value, str):
        raise ValueError('text values must be strings')
    return value


def _color(value):
    if not isinstance(value, str) or len(value) != 7 or not value.startswith('#'):
        raise ValueError('colors must be #RRGGBB')
    return RGBColor.from_string(value[1:])


def _style(frame, text, points, *, bold=False, color='#000000'):
    frame.clear()
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = frame.margin_right = Inches(.08)
    frame.margin_top = frame.margin_bottom = Inches(.04)
    for index, line in enumerate(text.split('\n')):
        p = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        run = p.add_run()
        run.text = line
        run.font.name = 'Meiryo'
        run.font.size = Pt(points)
        run.font.bold = bold
        run.font.color.rgb = _color(color)


def _shape(slide, kind, rect, name, *, fill='#FFFFFF', line='#000000'):
    shape = slide.shapes.add_shape(kind, *(Inches(v) for v in rect))
    shape.name = name
    shape.fill.solid()
    shape.fill.fore_color.rgb = _color(fill)
    shape.line.color.rgb = _color(line)
    shape.line.width = Pt(1)
    return shape


def _textbox(slide, rect, text, name, points, *, bold=False):
    shape = slide.shapes.add_textbox(*(Inches(v) for v in rect))
    shape.name = name
    _style(shape.text_frame, text, points, bold=bold)
    return shape


def enrich(source, patterns, *, destination=None, dry_run=False):
    """Append only; copy every non-target ZIP member without modification.

    Coordinates are inches. Callers choose a free region; overlap with existing
    objects is not resolved. Dry-run builds the same native objects in memory.
    """
    if not isinstance(patterns, list) or not patterns:
        raise ValueError('at least one pattern is required')
    if not dry_run and (destination is None or Path(source).resolve() == Path(destination).resolve() or os.path.lexists(destination)):
        raise ValueError('destination must be a new file distinct from source')
    prs = Presentation(str(source))
    changed, report = {}, []
    for index, spec in enumerate(patterns):
        if not isinstance(spec, dict):
            raise ValueError('pattern must be a mapping')
        mode, number = spec.get('pattern'), spec.get('slide')
        if mode not in ('cards', 'regions', 'agenda', 'flow', 'table'):
            raise ValueError('unknown pattern')
        if type(number) is not int or not 1 <= number <= len(prs.slides):
            raise ValueError('slide must be a valid 1-based slide number')
        rect = [_number(spec.get(k), k, positive=k in ('width','height')) for k in ('x','y','width','height')]
        x,y,w,h = rect
        if Inches(x+w) > prs.slide_width or Inches(y+h) > prs.slide_height:
            raise ValueError('pattern extends beyond slide bounds')
        gap = _number(spec.get('gap', .2), 'gap')
        points = _number(spec.get('font_size', 14), 'font_size', positive=True)
        if not 6 <= points <= 72:
            raise ValueError('font_size must be between 6 and 72')
        accent = spec.get('accent', '#000000')
        _color(accent)
        slide = prs.slides[number-1]
        before = len(slide.shapes)
        prefix = f'enrich-{index+1}-{mode}'
        if mode == 'table':
            if 'csv' in spec and 'rows' in spec:
                raise ValueError('use either csv or rows, not both')
            if 'csv' in spec:
                with Path(_text(spec['csv'])).open(encoding='utf-8-sig', newline='') as stream:
                    rows = list(csv.reader(stream))
            else:
                rows = spec.get('rows')
            header_fill = spec.get('header_fill', '#222222')
            header_text = spec.get('header_text', '#FFFFFF')
            row_fills = spec.get('row_fills', ['#FFFFFF', '#F2F2F2'])
            border = spec.get('border_color', '#B3B3B3')
            if not isinstance(row_fills, list) or not row_fills:
                raise ValueError('row_fills must be a nonempty list of colors')
            for value in [header_fill, header_text, border, *row_fills]:
                _color(value)
            if (not isinstance(rows,list) or not 2 <= len(rows) <= 30 or not isinstance(rows[0],list)
                    or not 1 <= len(rows[0]) <= 12 or any(not isinstance(r,list) or len(r)!=len(rows[0]) for r in rows)):
                raise ValueError('table requires 2..30 rectangular rows with 1..12 columns')
            for row in rows:
                for value in row:
                    _text(value)
            shape = slide.shapes.add_table(len(rows),len(rows[0]),*(Inches(v) for v in rect))
            shape.name = prefix
            table = shape.table
            for ri,row in enumerate(rows):
                for ci,value in enumerate(row):
                    cell = table.cell(ri,ci)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = _color(header_fill if ri==0 else row_fills[(ri-1)%len(row_fills)])
                    _style(cell.text_frame,value,points,bold=ri==0,color=header_text if ri==0 else '#000000')
                    props = cell._tc.get_or_add_tcPr()
                    for side in ('L', 'R', 'T', 'B'):
                        tag = 'a:ln'+side
                        line = OxmlElement(tag)
                        line.set('w', '6350') # 0.5 points
                        fill = OxmlElement('a:solidFill')
                        rgb = OxmlElement('a:srgbClr'); rgb.set('val', border[1:].upper())
                        fill.append(rgb); line.append(fill)
                        props.append(line)
        else:
            items = spec.get('items')
            if not isinstance(items,list) or not 1 <= len(items) <= 24:
                raise ValueError('items must contain 1..24 mappings')
            for item in items:
                if not isinstance(item,dict):
                    raise ValueError('item must be a mapping')
                _text(item.get('title',''))
                _text(item.get('body',''))
                _text(item.get('label',''))
            n = len(items)
            columns = spec.get('columns', n)
            if mode in ('regions','flow'):
                columns = n
            elif mode == 'agenda':
                columns = 1
            if type(columns) is not int or not 1 <= columns <= n:
                raise ValueError('columns must be between 1 and item count')
            row_count = (n+columns-1)//columns
            cw,ch = (w-gap*(columns-1))/columns,(h-gap*(row_count-1))/row_count
            if cw < .7 or ch < .7:
                raise ValueError('cells must be at least 0.7 inches wide and high')
            if mode == 'flow' and n>1 and gap < .15:
                raise ValueError('flow gap must be at least 0.15 inches')
            for i,item in enumerate(items):
                cx,cy = x+(i%columns)*(cw+gap), y+(i//columns)*(ch+gap)
                name = f'{prefix}-{i+1}'
                kind = MSO_SHAPE.ROUNDED_RECTANGLE if mode in ('regions','agenda') else MSO_SHAPE.RECTANGLE
                _shape(slide,kind,(cx,cy,cw,ch),name+'-box',line=accent)
                if mode == 'agenda':
                    _textbox(slide,(cx+.05,cy+.03,cw-.1,ch-.06),item.get('title','')+'\n'+item.get('body','')+('  '+item['label'] if item.get('label') else ''),name+'-text',points,bold=False)
                else:
                    header = min(.45,ch*.3)
                    if mode == 'regions':
                        badge = _shape(slide,MSO_SHAPE.ROUNDED_RECTANGLE,(cx+.05,cy+.03,cw-.1,header),name+'-label',fill=accent,line=accent)
                        _style(badge.text_frame,item.get('title',''),points,bold=True,color='#FFFFFF')
                    else:
                        _shape(slide,MSO_SHAPE.RECTANGLE,(cx,cy,cw,.035),name+'-bar',fill=accent,line=accent)
                        _textbox(slide,(cx+.04,cy+.04,cw-.08,header),item.get('title',''),name+'-title',points,bold=True)
                    _textbox(slide,(cx+.04,cy+header+.07,cw-.08,ch-header-.1),item.get('body',''),name+'-body',points)
                if mode == 'flow' and i<n-1:
                    arrow = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                        Inches(cx+cw+gap*.15), Inches(cy+ch/2),
                        Inches(cx+cw+gap*.85), Inches(cy+ch/2))
                    arrow.name = name+'-arrow'
                    arrow.line.color.rgb = _color(accent)
                    arrow.line.width = Pt(1.5)
                    end = OxmlElement('a:tailEnd')
                    end.set('type', 'triangle')
                    end.set('w', 'med'); end.set('len', 'med')
                    arrow.line._get_or_add_ln().append(end)
        member = str(slide.part.partname).lstrip('/')
        changed[member] = slide.part.blob
        report.append(dict(pattern=mode,slide=number,added_shapes=len(slide.shapes)-before,
                           shape_ids=[s.shape_id for s in list(slide.shapes)[before:]]))
    with ZipFile(source) as original:
        if len(original.namelist()) != len(set(original.namelist())):
            raise ValueError('duplicate ZIP members unsupported')
        if not dry_run:
            dst = Path(destination)
            with tempfile.TemporaryDirectory(prefix='.pptx-enrich-', dir=dst.parent) as staging:
                staged = Path(staging)/'complete.pptx'
                with ZipFile(staged,'w') as output:
                    output.comment = original.comment
                    for entry in original.infolist():
                        output.writestr(entry, changed.get(entry.filename, original.read(entry.filename)))
                Presentation(staged)
                os.link(staged,dst)
    return dict(dry_run=dry_run,patterns=report,added_shapes=sum(r['added_shapes'] for r in report))


def load_patterns(path):
    """Read TOML and resolve CSV references relative to its directory."""
    path = Path(path)
    with path.open('rb') as stream:
        patterns = tomllib.load(stream).get('patterns')
    if not isinstance(patterns, list) or not patterns:
        raise ValueError('at least one pattern is required')
    for spec in patterns:
        if isinstance(spec, dict) and isinstance(spec.get('csv'), str):
            spec['csv'] = str(path.parent / spec['csv'])
    return patterns


def convert(source, output, *, marp="marp", allow_local_files=False):
    source, output = Path(source), Path(output)
    if source.suffix.lower() != ".md":
        raise ValueError("input must be a .md file")
    if output.suffix.lower() != ".pptx":
        raise ValueError("output must be a .pptx file")
    if source.resolve() == output.resolve() or output.exists():
        raise ValueError("output must be a new file distinct from input")
    executable = shutil.which(marp)
    if executable is None:
        raise RuntimeError("Marp CLI not installed; install and pin @marp-team/marp-cli separately")
    command = [executable, str(source), "--pptx", "--output", str(output)]
    if allow_local_files:
        command.append("--allow-local-files")
    result = subprocess.run(command, capture_output=True, text=True, check=False, timeout=120)
    if result.returncode != 0:
        raise RuntimeError(f"Marp conversion failed (exit {result.returncode}): {result.stderr.strip()}")
    if not output.is_file():
        raise RuntimeError("Marp returned success but produced no PPTX")
    try:
        slides = len(Presentation(output).slides)
    except Exception as exc:
        raise RuntimeError("Marp output is not a readable PPTX") from exc
    return {"source": str(source), "output": str(output), "slides": slides, "backend": "marp"}


def read_markdown_input(*, source=None, text=None, stdin=None):
    """Resolve Markdown from file, inline text, or stdin without altering CLI semantics."""
    if (source is None) == (text is None):
        raise ValueError("provide exactly one Markdown source")
    if text is not None:
        return text, "inline"
    if source == "-":
        return (sys.stdin if stdin is None else stdin).read(), "stdin"
    return Path(source).read_text(encoding="utf-8"), str(source)



def run(*, source=None, text=None, output, marp="marp", mmdc=None, background="#FFFFFF", svg_foreground=None, enrich_patterns=None):
    output = Path(output)
    if (source is None) == (text is None):
        raise ValueError("provide exactly one Markdown source")
    if enrich_patterns is not None:
        if output.suffix.lower() != ".pptx" or os.path.lexists(output):
            raise ValueError("output must be a new .pptx file")
        # Keep the Marp intermediate private until all native edits succeed.
        with tempfile.TemporaryDirectory(prefix="ironmate-enrich-", dir=output.parent) as directory:
            intermediate = Path(directory) / "marp.pptx"
            result = run(source=source, text=text, output=intermediate, marp=marp,
                         mmdc=mmdc, background=background, svg_foreground=svg_foreground)
            edit_report = enrich(intermediate, enrich_patterns, destination=output)
        result["conversion"]["output"] = str(output)
        result["enrichment"] = edit_report
        result["inspection"] = inspect(output)
        return result
    if mmdc is not None:
        raw, _ = read_markdown_input(source=source, text=text)
        if not raw.strip():
            raise ValueError("Markdown must not be empty")
        with tempfile.TemporaryDirectory(prefix="ironmate-svg-") as directory:
            working = Path(directory)
            prepared, images = render_mermaid(raw, working, mmdc=mmdc, background=background, foreground=svg_foreground)
            input_path = working / "input.md"
            input_path.write_text(prepared, encoding="utf-8")
            conversion = convert(input_path, output, marp=marp, allow_local_files=True)
            conversion["mermaid_svg_count"] = len(images)
    elif text is not None:
        if not text.strip():
            raise ValueError("Markdown must not be empty")
        with tempfile.TemporaryDirectory(prefix="ironmate-md-") as directory:
            input_path = Path(directory) / "input.md"
            input_path.write_text(text, encoding="utf-8")
            conversion = convert(input_path, output, marp=marp)
            conversion["source"] = "inline"
    else:
        conversion = convert(source, output, marp=marp)
    report = inspect(output)
    return {"conversion": conversion, "inspection": report}


def review(*, source=None, text=None, directory, marp='marp', mmdc=None,
           background='#FFFFFF', svg_foreground=None, enrich_patterns=None,
           sha=None, run_id=None):
    """One-command source preview + final PPTX structural inspection.

    Never claims that Marp HTML/PDF includes later native PPTX edits.
    Only a completed bundle publishes final.pptx; existing paths are preserved.
    """
    import base64
    import hashlib
    import html as html_escape
    import uuid
    from pptx_preview import render_preview_html
    target = Path(directory)
    if os.path.lexists(target) or not target.parent.is_dir():
        raise ValueError('review directory must be new with an existing parent')
    raw, _ = read_markdown_input(source=source, text=text)
    if not raw.strip():
        raise ValueError('Markdown must not be empty')
    executable = shutil.which(marp)
    if executable is None:
        raise RuntimeError('Marp CLI not installed')
    if sha is None:
        sha = subprocess.run(['git', '-C', str(Path(__file__).resolve().parent), 'rev-parse', 'HEAD'],
                             capture_output=True, text=True, check=True, timeout=10).stdout.strip()
    if len(sha) != 40 or any(c not in '0123456789abcdef' for c in sha):
        raise ValueError('review SHA must be a full lowercase Git SHA')
    run_id = run_id or uuid.uuid4().hex
    target.mkdir()  # Exclusive, before any render work.
    report = {'stage': 'source', 'source_sha256': hashlib.sha256(raw.encode()).hexdigest(),
              'tested_sha': sha, 'run_id': run_id, 'is_final_pptx_render': False,
              'native_pptx_edits_reflected': False,
              'scope': 'Marp source preview + separate final PPTX structural inspection'}
    try:
        # Keep the final PPTX private until previews and PDF validation succeed.
        with tempfile.TemporaryDirectory(prefix='.review-', dir=target) as scratch:
            private = Path(scratch) / 'final.pptx'
            prepared = raw
            if mmdc is not None:
                prepared, images = render_mermaid(raw, target, mmdc=mmdc,
                                                 background=background, foreground=svg_foreground)
                # Keep SVG previews portable instead of retaining temporary paths.
                for image in images:
                    uri = 'data:image/svg+xml;base64,' + base64.b64encode(image.read_bytes()).decode()
                    prepared = prepared.replace('./' + image.name, uri)
            md = target / 'source.md'
            md.write_text(prepared, encoding='utf-8')
            report['stage'] = 'pptx'
            # Keep the existing PPTX path's original asset-resolution behavior.
            result = run(source=source, text=text, output=private, marp=marp, mmdc=mmdc,
                         background=background, svg_foreground=svg_foreground, enrich_patterns=enrich_patterns)
            report['stage'] = 'html'
            html = target / 'preview.html'
            proc = subprocess.run([executable, str(md), '--template', 'bare', '-o', str(html)],
                                  capture_output=True, text=True, timeout=120, check=False)
            if proc.returncode != 0 or not html.is_file():
                raise RuntimeError('Marp HTML conversion failed')
            # Resolve ordinary relative assets against the original source folder.
            # HTML portability still depends on those assets; PDF/PNG are standalone.
            base = (Path(source).resolve().parent if source is not None else Path.cwd()).as_uri() + '/'
            markup = html.read_text(encoding='utf-8')
            if '<head>' not in markup:
                raise ValueError('Marp output has no HTML head')
            markup = markup.replace('<head>', '<head><base href="' + html_escape.escape(base, quote=True) + '">', 1)
            html.write_text(markup, encoding='utf-8')
            report['stage'] = 'preview'
            preview = render_preview_html(html, target, expected_slides=len(result['inspection']['slides']),
                                          sha=sha, run_id=run_id)
            if hashlib.sha256((Path(source).read_text(encoding='utf-8') if source is not None else text).encode()).hexdigest() != report['source_sha256']:
                raise ValueError('Markdown source changed during review')
            final = target / 'final.pptx'
            result['conversion']['output'] = str(final)
            result['inspection']['source'] = str(final)
            (target / 'inspection.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
            final_hash = hashlib.sha256(private.read_bytes()).hexdigest()
            os.link(private, final)
            report.update(stage='complete', preview=preview, final_pptx_sha256=final_hash)
    except Exception as exc:
        report.update(failed_stage=report['stage'], stage='failed', error_type=type(exc).__name__)
        raise RuntimeError(f"review failed at {report['failed_stage']} ({type(exc).__name__}); see review.json") from None
    finally:
        (target / 'review.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return report


def table_pattern(path, *, pattern="table", slide=1, region=(.5, 1.2, 9, 4), columns=None):
    """Turn CSV or the first live Markdown table into a native table/diagram.

    Diagram mapping: first cell is title; remaining cells become header: value
    body lines. No inference, scoring, or new Markdown parser.
    """
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            data = list(csv.reader(stream))
        if len(data) < 2 or not data[0] or any(len(row) != len(data[0]) for row in data):
            raise ValueError("CSV requires header and rectangular body rows")
        headers, rows = data[0], data[1:]
    elif path.suffix.lower() == ".md":
        headers, rows = markdown_table_to_rows(path.read_text(encoding="utf-8-sig"))
        if not headers or not rows:
            raise ValueError("Markdown must contain a live GFM table with body rows")
    else:
        raise ValueError("table input must be .csv or .md")
    if pattern not in ("table", "cards", "regions", "flow"):
        raise ValueError("table pattern must be table, cards, regions or flow")
    spec = dict(pattern=pattern, slide=slide, **dict(zip(("x","y","width","height"),region)))
    if pattern == "table":
        spec["rows"] = [headers, *rows]
    else:
        spec["items"] = [dict(title=row[0], body="\n".join(f"{key}: {value}" for key,value in zip(headers[1:],row[1:]))) for row in rows]
        if columns is not None:
            spec["columns"] = columns
    return spec


def main(argv=None):
    """Single supported Markdown-to-PPTX CLI, with optional SVG preprocessing."""
    parser = make_parser([
        Argument(("source",), {"nargs": "?", "help": "Markdown file path or - for stdin"}),
        Argument(("--text",), {"help": "Inline Markdown source"}),
        Argument(("-o", "--output"), {"type": Path}),
        Argument(("--review",), {"type": Path, "help": "New directory: final PPTX inspection plus Marp HTML/PDF/PNG preview"}),
        Argument(("--review-sha",), {"help": "Tested Git SHA; defaults to repository HEAD"}),
        Argument(("--review-run-id",), {}),
        Argument(("--marp",), {"default": "marp"}),
        Argument(("--mmdc",), {"help": "Optional Mermaid SVG pre-renderer"}),
        Argument(("--background",), {"default": "#FFFFFF"}),
        Argument(("--svg-foreground",), {}),
        Argument(("--table",), {"type": Path, "help": "CSV or Markdown table; TOML not required"}),
        Argument(("--table-pattern",), {"choices": ("table","cards","regions","flow"), "default": "table"}),
        Argument(("--table-slide",), {"type": int, "default": 1}),
        Argument(("--table-region",), {"type": float, "nargs": 4, "default": (.5,1.2,9,4), "metavar": ("X","Y","WIDTH","HEIGHT")}),
        Argument(("--table-columns",), {"type": int}),
        Argument(("--enrich",), {"type": Path, "help": "TOML patterns applied with python-pptx after Marp"}),
    ], description=__doc__)
    args = parser.parse_args(argv)
    if (args.source is None) == (args.text is None):
        parser.error("provide exactly one of source or --text")
    if (args.output is None) == (args.review is None):
        parser.error("provide exactly one of --output or --review")
    if args.review is None and (args.review_sha is not None or args.review_run_id is not None):
        parser.error("review metadata requires --review")
    inline = args.text if args.text is not None else sys.stdin.read() if args.source == "-" else None
    if inline is not None and not inline.strip():
        parser.error("Markdown content must not be empty")
    if args.table is not None and args.enrich is not None:
        parser.error("use --table or optional --enrich configuration, not both")
    patterns = None
    if args.table is not None:
        patterns = [table_pattern(args.table, pattern=args.table_pattern, slide=args.table_slide,
                                  region=args.table_region, columns=args.table_columns)]
    if args.enrich is not None:
        patterns = load_patterns(args.enrich)
    common = dict(source=None if inline is not None else args.source, text=inline,
                  marp=args.marp, mmdc=args.mmdc, background=args.background,
                  svg_foreground=args.svg_foreground, enrich_patterns=patterns)
    if args.review is not None:
        result = review(directory=args.review, sha=args.review_sha, run_id=args.review_run_id, **common)
    else:
        result = run(output=args.output, **common)
    if args.source == "-" and args.review is None:
        result["conversion"]["source"] = "stdin"
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
