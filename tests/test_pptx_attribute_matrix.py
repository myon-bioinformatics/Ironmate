"""Attribute insertion and narrow Open XML mutation experiments."""
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from pptx_replace import transform

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def make_input(path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    shape.text_frame.paragraphs[0].add_run().text = "Test"
    presentation.save(path)


def properties(path):
    with ZipFile(path) as archive:
        xml = ET.fromstring(archive.read("ppt/slides/slide1.xml"))
    return xml.find(".//a:rPr", NS)


@pytest.mark.parametrize("operation,rule,expected", [
    ("bold", {"value": True}, {"b": "1"}),
    ("bold", {"value": False}, {"b": "0"}),
    ("font_size", {"points": 18}, {"sz": "1800"}),
    ("text_color", {"color": "#123456"}, {"rgb": "123456"}),
])
def test_absent_attribute_insertion(tmp_path, operation, rule, expected):
    source, output = tmp_path / "source.pptx", tmp_path / "output.pptx"
    make_input(source)
    before = properties(source)
    assert before is None or not any(key in before.attrib for key in ("b", "sz"))
    if operation == "text_color":
        # Current CLI deliberately rejects inherited/theme color; record unsupported.
        with pytest.raises(ValueError, match="explicit RGB"):
            transform(source, [{"operation": operation, **rule}], destination=output)
        assert not output.exists()
        return
    assert transform(source, [{"operation": operation, **rule}], destination=output)["change_count"] == 1
    after = properties(output)
    assert after is not None
    for key, value in expected.items():
        assert after.attrib[key] == value
    assert Presentation(output).slides[0].shapes[0].text == "Test"


def test_sequential_style_edits_and_dry_run(tmp_path):
    source, output = tmp_path / "source.pptx", tmp_path / "styled.pptx"
    make_input(source)
    rules = [
        {"operation": "font_size", "points": 18},
        {"operation": "bold", "value": True},
        {"operation": "bold", "value": False},
    ]
    preview = transform(source, rules, dry_run=True)
    actual = transform(source, rules, destination=output)
    assert preview["change_count"] == actual["change_count"] == 3
    run = Presentation(output).slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.size == Pt(18)
    assert run.font.bold is False
    assert properties(output).attrib["b"] == "0"


def test_xml_fallback_inserts_explicit_rgb_without_changing_other_members(tmp_path):
    """Experimental XML-only fixture, not a production CLI mutation."""
    source, output = tmp_path / "source.pptx", tmp_path / "xml-color.pptx"
    make_input(source)
    uri = NS["a"]
    with ZipFile(source) as archive:
        names = archive.namelist()
        members = {name: archive.read(name) for name in names}
    root = ET.fromstring(members["ppt/slides/slide1.xml"])
    run = root.find(".//a:r", NS)
    assert run is not None
    rpr = run.find("a:rPr", NS)
    if rpr is None:
        rpr = ET.Element(f"{{{uri}}}rPr")
        run.insert(0, rpr)
    fill = ET.SubElement(rpr, f"{{{uri}}}solidFill")
    ET.SubElement(fill, f"{{{uri}}}srgbClr", {"val": "123456"})
    members["ppt/slides/slide1.xml"] = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with ZipFile(output, "x") as archive:
        for name in names:
            archive.writestr(name, members[name])
    with ZipFile(source) as before, ZipFile(output) as after:
        assert set(before.namelist()) == set(after.namelist())
        assert all(before.read(name) == after.read(name) for name in names if name != "ppt/slides/slide1.xml")
    run = Presentation(output).slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert run.font.color.rgb == RGBColor(18, 52, 86)
    assert run.text == "Test"
