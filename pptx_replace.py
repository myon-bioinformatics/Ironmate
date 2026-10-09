"""Conservative PPTX attribute replacement prototype. Python 3.11+."""
import argparse
import json
from pathlib import Path
import tomllib

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt


def _rgb(value):
    if not isinstance(value, str) or len(value) != 7 or value[0] != "#":
        raise ValueError("color must be #RRGGBB")
    try:
        return RGBColor.from_string(value[1:])
    except ValueError as exc:
        raise ValueError("color must be #RRGGBB") from exc


def transform(source, rules, *, dry_run=False, destination=None):
    if not isinstance(rules, list) or not rules:
        raise ValueError("at least one rule is required")
    if destination is not None and not dry_run:
        src, dst = Path(source), Path(destination)
        if src.resolve() == dst.resolve() or dst.exists():
            raise ValueError("destination must be a new file distinct from source")
    presentation = Presentation(str(source))
    changes = []
    for index, rule in enumerate(rules):
        operation = rule.get("operation")
        if operation not in {"replace_text", "text_color", "slide_background", "font_size", "bold"}:
            raise ValueError(f"unsupported operation: {operation!r}")
        slide_number = rule.get("slide")
        if slide_number is not None and (type(slide_number) is not int or not 1 <= slide_number <= len(presentation.slides)):
            raise ValueError("slide must be a valid 1-based slide number")
        if operation == "replace_text":
            old, new = rule.get("old"), rule.get("new")
            if not isinstance(old, str) or not old or not isinstance(new, str):
                raise ValueError("replace_text requires nonempty old and string new")
        elif operation in {"text_color", "slide_background"}:
            color = _rgb(rule.get("color"))
        elif operation == "font_size":
            size = rule.get("points")
            if type(size) not in (int, float) or not 1 <= size <= 400:
                raise ValueError("font_size requires numeric points between 1 and 400")
        elif operation == "bold":
            if type(rule.get("value")) is not bool:
                raise ValueError("bold requires a boolean value")
        for number, slide in enumerate(presentation.slides, 1):
            if slide_number is not None and number != slide_number:
                continue
            if operation == "slide_background":
                # Only explicit solid RGB backgrounds are supported.
                fill = slide.background.fill
                if fill.type is None:
                    raise ValueError("inherited slide background is unsupported")
                from pptx.enum.dml import MSO_FILL_TYPE
                from pptx.dml.color import MSO_COLOR_TYPE
                if fill.type != MSO_FILL_TYPE.SOLID or fill.fore_color.type != MSO_COLOR_TYPE.RGB:
                    raise ValueError("only explicit solid RGB slide backgrounds are supported")
                if fill.fore_color.rgb != color:
                    changes.append({"rule": index, "slide": number, "operation": operation})
                    if not dry_run:
                        fill.fore_color.rgb = color
                continue
            for shape_index, shape in enumerate(slide.shapes, 1):
                if not shape.has_text_frame:
                    continue
                for paragraph_index, paragraph in enumerate(shape.text_frame.paragraphs, 1):
                    for run_index, run in enumerate(paragraph.runs, 1):
                        if operation == "replace_text":
                            count = run.text.count(old)
                            if count:
                                changes.append({"rule": index, "slide": number, "shape": shape_index, "paragraph": paragraph_index, "run": run_index, "count": count})
                                if not dry_run:
                                    run.text = run.text.replace(old, new)
                        elif operation == "font_size":
                            if run.font.size != Pt(size):
                                changes.append({"rule": index, "slide": number, "shape": shape_index, "paragraph": paragraph_index, "run": run_index, "count": 1})
                                if not dry_run:
                                    run.font.size = Pt(size)
                        elif operation == "bold":
                            if run.font.bold is not rule["value"]:
                                changes.append({"rule": index, "slide": number, "shape": shape_index, "paragraph": paragraph_index, "run": run_index, "count": 1})
                                if not dry_run:
                                    run.font.bold = rule["value"]
                        elif operation == "text_color":
                            from pptx.dml.color import MSO_COLOR_TYPE
                            if run.font.color.type != MSO_COLOR_TYPE.RGB:
                                raise ValueError("text_color requires explicit RGB run colors")
                            if run.font.color.rgb != color:
                                changes.append({"rule": index, "slide": number, "shape": shape_index, "paragraph": paragraph_index, "run": run_index, "count": 1})
                                if not dry_run:
                                    run.font.color.rgb = color
    if not dry_run:
        if destination is None:
            raise ValueError("destination is required unless dry-run")
        dst = Path(destination)
        try:
            with dst.open("xb") as output:
                presentation.save(output)
        except BaseException:
            dst.unlink(missing_ok=True)
            raise
    return {"dry_run": dry_run, "changes": changes, "change_count": sum(c.get("count", 1) for c in changes)}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--rules", type=Path)
    parser.add_argument("--replace", nargs=2, metavar=("OLD", "NEW"))
    parser.add_argument("--text-color")
    parser.add_argument("--background")
    parser.add_argument("--font-size", type=float)
    parser.add_argument("--bold", choices=("true", "false"))
    parser.add_argument("--output", type=Path)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args(argv)
    if not args.dry_run and args.output is None:
        parser.error("--output is required unless --dry-run")
    if args.rules and any((args.replace, args.text_color, args.background, args.font_size is not None, args.bold is not None)):
        parser.error("--rules cannot be combined with inline operations")
    if args.rules:
        with args.rules.open("rb") as stream:
            config = tomllib.load(stream)
        rules = config.get("rules")
    else:
        rules = []
        if args.replace:
            rules.append({"operation": "replace_text", "old": args.replace[0], "new": args.replace[1]})
        if args.text_color:
            rules.append({"operation": "text_color", "color": args.text_color})
        if args.background:
            rules.append({"operation": "slide_background", "color": args.background})
        if args.font_size is not None:
            rules.append({"operation": "font_size", "points": args.font_size})
        if args.bold is not None:
            rules.append({"operation": "bold", "value": args.bold == "true"})
        if not rules:
            parser.error("specify --rules or an inline operation")
    result = transform(args.source, rules, dry_run=args.dry_run, destination=args.output)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
