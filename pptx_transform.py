"""Conservative PPTX attribute replacement prototype (Python >=3.11)."""
import argparse
import json
import sys
import tomllib
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor


def _color(value):
    if not isinstance(value, str) or len(value) != 7 or not value.startswith("#"):
        raise ValueError("color must be #RRGGBB")
    return RGBColor.from_string(value[1:])


def transform(source, rules, output=None, dry_run=False):
    prs = Presentation(str(source))
    changes = []
    for idx, rule in enumerate(rules):
        kind = rule["operation"]
        slide_no = rule.get("slide")
        if slide_no is not None and (not isinstance(slide_no, int) or isinstance(slide_no, bool) or not 1 <= slide_no <= len(prs.slides)):
            raise ValueError("invalid slide number")
        if kind not in {"replace_text", "font_color", "slide_background"}:
            raise ValueError(f"unsupported operation: {kind}")
        if kind == "replace_text" and (not isinstance(rule.get("old"), str) or not rule["old"] or not isinstance(rule.get("new"), str)):
            raise ValueError("replace_text requires nonempty old and string new")
        if kind in {"font_color", "slide_background"}:
            color = _color(rule["new"])
        for si, slide in enumerate(prs.slides, 1):
            if slide_no is not None and si != slide_no:
                continue
            if kind == "slide_background":
                # Initial scope: direct slide background only, no theme/master resolution.
                if rule.get("old") is not None:
                    raise ValueError("background precondition not supported yet")
                changes.append({"rule": idx, "slide": si, "operation": kind})
                if not dry_run:
                    slide.background.fill.solid()
                    slide.background.fill.fore_color.rgb = color
                continue
            for shape_i, shape in enumerate(slide.shapes, 1):
                if not shape.has_text_frame:
                    continue
                for pi, para in enumerate(shape.text_frame.paragraphs, 1):
                    for ri, run in enumerate(para.runs, 1):
                        if kind == "replace_text":
                            count = run.text.count(rule["old"])
                            if not count:
                                continue
                            changes.append({"rule": idx, "slide": si, "shape": shape_i, "paragraph": pi, "run": ri, "count": count})
                            if not dry_run:
                                run.text = run.text.replace(rule["old"], rule["new"])
                        else:
                            if not run.text:
                                continue
                            if rule.get("old") is not None:
                                old = _color(rule["old"])
                                if run.font.color.type is None or run.font.color.type != 1 or run.font.color.rgb != old:
                                    continue
                            changes.append({"rule": idx, "slide": si, "shape": shape_i, "paragraph": pi, "run": ri, "count": 1})
                            if not dry_run:
                                run.font.color.rgb = color
    if not dry_run:
        if output is None or Path(source).resolve() == Path(output).resolve():
            raise ValueError("output must be distinct from input")
        prs.save(str(output))
        Presentation(str(output))
    return changes


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--rules", required=True, type=Path)
    parser.add_argument("--output", type=Path)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args(argv)
    try:
        with args.rules.open("rb") as f:
            config = tomllib.load(f)
        rules = config["rules"]
        if not isinstance(rules, list) or not rules:
            raise ValueError("rules must be a nonempty list")
        if not args.dry_run and args.output is None:
            raise ValueError("--output required without --dry-run")
        changes = transform(args.input, rules, args.output, args.dry_run)
        print(json.dumps({"dry_run": args.dry_run, "changes": changes, "count": sum(x.get("count", 1) for x in changes)}, ensure_ascii=False))
        return 0
    except (ValueError, KeyError, OSError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
