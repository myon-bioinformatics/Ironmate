"""Local Chromium fixtures prove capture plumbing, never PowerPoint rendering."""
from contextlib import contextmanager
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import json
import os
from pathlib import Path
import subprocess
import sys
import threading
from types import SimpleNamespace

from pptx import Presentation
import pytest

import pptx_web_capture as capture


@pytest.fixture
def source(tmp_path):
    prs = Presentation()
    for text in ('Alpha', 'Beta'):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_textbox(0, 0, 1000000, 1000000).text = text
    path = tmp_path / 'synthetic.pptx'
    prs.save(path)
    return path


def plan():
    return {'document': '#document', 'document_text': 'synthetic.pptx',
            'surface': '#slide', 'frames': ['iframe'], 'blockers': ['#signin'], 'slides': [
                {'number': 1, 'select': '#one', 'ready': '#ready', 'ready_text': 'Alpha'},
                {'number': 2, 'select': '#two', 'ready': '#ready', 'ready_text': 'Beta'}]}


def args(source, tmp_path):
    return dict(source=source, plan=plan(), url='http://127.0.0.1:1234/view',
                output=tmp_path / 'capture', kit=tmp_path / 'kit', sha='a' * 40,
                run_id='unit', attempt='1', fixture=True, cdp=None, timeout=500)


@pytest.mark.parametrize('patch', [
    {'url': 'https://evil.example/view'}, {'url': 'http://127.0.0.1.evil.example/'},
    {'url': 'http://name:secret@127.0.0.1/'}, {'sha': 'short'}, {'run_id': ''},
    {'attempt': ''}, {'timeout': 0}, {'timeout': True}, {'timeout': 60001},
    {'fixture': False, 'url': 'https://powerpoint.cloud.microsoft/view', 'cdp': None},
    {'fixture': False, 'url': 'https://powerpoint.cloud.microsoft/view', 'cdp': 'http://remote:9222'},
    {'fixture': True, 'cdp': 'http://localhost:9222'},
])
def test_preflight_rejects_before_any_network_or_output(source, tmp_path, monkeypatch, patch):
    config = args(source, tmp_path)
    config.update(patch)
    monkeypatch.setattr(capture, 'verify_kit', lambda *_: pytest.fail('reached kit after invalid input'))
    with pytest.raises(ValueError):
        capture.preflight(**config)
    assert not config['output'].exists()


@pytest.mark.parametrize('patch', [
    {'document_text': 'another.pptx'}, {'document': ''}, {'surface': None},
    {'frames': 'iframe'}, {'blockers': []}, {'slides': []},
    {'slides': [{'number': 3, 'select': '#one', 'ready': '#ready', 'ready_text': 'X'}]},
    {'slides': [plan()['slides'][0], plan()['slides'][0]]},
])
def test_plan_rejection(source, tmp_path, monkeypatch, patch):
    config = args(source, tmp_path)
    config['plan'].update(patch)
    monkeypatch.setattr(capture, 'verify_kit', lambda *_: pytest.fail('kit must not run'))
    with pytest.raises(ValueError):
        capture.preflight(**config)
    assert not config['output'].exists()


def test_existing_output_and_symlink_unchanged(source, tmp_path, monkeypatch):
    config = args(source, tmp_path)
    target = config['output']
    target.mkdir()
    (target / 'sentinel').write_text('old evidence')
    with pytest.raises(ValueError):
        capture.preflight(**config)
    assert (target / 'sentinel').read_text() == 'old evidence'
    link = tmp_path / 'dangling'
    link.symlink_to(tmp_path / 'missing')
    config['output'] = link
    with pytest.raises(ValueError):
        capture.preflight(**config)
    assert link.is_symlink()


def test_launch_failure_is_never_complete(source, tmp_path, monkeypatch):
    config = args(source, tmp_path)
    monkeypatch.setattr(capture, 'verify_kit', lambda *_: tmp_path)
    stages = []
    monkeypatch.setattr(capture, 'seal', lambda *a, **kw: stages.append(kw['stage']))
    @contextmanager
    def unavailable():
        raise RuntimeError('https://private.example/?token=SECRET')
        yield
    monkeypatch.setitem(sys.modules, 'playwright.sync_api', SimpleNamespace(
        sync_playwright=unavailable, expect=None))
    with pytest.raises(RuntimeError, match='launch'):
        capture.capture(**config)
    report = json.loads((config['output'] / 'capture.json').read_text())
    assert report['stage'] == 'failed' and report['failed_stage'] == 'launch'
    assert stages == ['failed']
    assert 'SECRET' not in json.dumps(report)


def test_evidence_tools_are_pinned_and_unmodified(tmp_path, monkeypatch):
    kit = tmp_path / 'kit'
    (kit / 'scripts').mkdir(parents=True)
    for name in capture.KIT_FILES:
        (kit / name).write_bytes(b'reviewed')
    def git(argv, **kwargs):
        return SimpleNamespace(stdout=(capture.KIT_SHA + '\n').encode() if 'rev-parse' in argv else b'reviewed')
    monkeypatch.setattr(capture.subprocess, 'run', git)
    assert capture.verify_kit(kit) == kit
    (kit / capture.KIT_FILES[0]).write_bytes(b'modified')
    with pytest.raises(ValueError, match='local changes'):
        capture.verify_kit(kit)


def test_seal_delegates_both_writer_and_validator(tmp_path, monkeypatch):
    commands = []
    monkeypatch.setattr(capture.subprocess, 'run', lambda argv, **kw: commands.append(argv))
    capture.seal(tmp_path, tmp_path / 'out', names=['slide-001.png'], project='chromium-fixture',
                 sha='a' * 40, run_id='run', attempt='2', stage='complete')
    assert len(commands) == 2
    assert all(argv[:2] == [sys.executable, '-S'] for argv in commands)
    assert commands[0][2].endswith('write_capture_evidence.py')
    assert commands[1][2].endswith('check_capture_evidence.py')
    assert all('--run-id' in cmd and '--run-attempt' in cmd and '--sha' in cmd for cmd in commands)


@contextmanager
def server():
    class Handler(BaseHTTPRequestHandler):
        def do_GET(self):
            if self.path == '/frame':
                body = '''<!doctype html><meta charset="utf-8">
<style>body{margin:0}#controls{height:32px;display:flex}button{height:32px}</style>
<div id="controls"><button id="one" onclick="show('Alpha','#ddd')">1</button>
<button id="two" onclick="show('Beta','#eee')">2</button></div>
<div id="slide" style="width:800px;height:450px;border:1px solid black;background:white">
<h1 id="ready">Loading</h1></div>
<script>function show(text,color) {document.querySelector('#ready').textContent=text;
document.querySelector('#slide').style.background=color;}</script>'''
            elif self.path == '/login':
                body = '<h1>Sign in</h1>'
            else:
                body = '<style>body{margin:0}h1{margin:0;height:64px;line-height:64px}iframe{position:absolute;left:0;top:64px;border:0}</style><h1 id="document">synthetic.pptx</h1><iframe src="/frame" style="width:1000px;height:650px"></iframe>'
                if self.path == '/overlay':
                    body += '<div id="signin" style="position:fixed;inset:0;background:white;z-index:100">Sign in</div>'
            self.send_response(200)
            self.send_header('Content-Type', 'text/html; charset=utf-8')
            self.end_headers()
            self.wfile.write(body.encode())
        def log_message(self, *args):
            pass
    http = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
    thread = threading.Thread(target=http.serve_forever, daemon=True)
    thread.start()
    try:
        yield f'http://127.0.0.1:{http.server_port}'
    finally:
        http.shutdown()
        http.server_close()
        thread.join()


@pytest.mark.skipif(os.environ.get('IRONMATE_RUN_WEB_CAPTURE') != '1',
                    reason='optional real Chromium capture lane not enabled')
@pytest.mark.parametrize('scenario', ['complete', 'login', 'wrong-slide', 'overlay'])
def test_real_chromium_and_shared_evidence(source, tmp_path, scenario):
    kit = Path(os.environ['IRONMATE_BTK_ROOT']).resolve()
    output = Path(os.environ.get('IRONMATE_WEB_EVIDENCE', str(tmp_path))) / scenario
    output.parent.mkdir(parents=True, exist_ok=True)
    options = args(source, tmp_path)
    options.update(kit=kit, output=output, timeout=1500,
                   sha=os.environ.get('IRONMATE_TEST_SHA', 'a' * 40),
                   run_id=os.environ.get('GITHUB_RUN_ID', 'local'),
                   attempt=os.environ.get('GITHUB_RUN_ATTEMPT', '1'))
    before = source.read_bytes()
    with server() as root:
        options['url'] = root + ('/' + scenario if scenario in ('login', 'overlay') else '/viewer')
        if scenario == 'wrong-slide':
            options['plan']['slides'][1]['ready_text'] = 'Never rendered'
        if scenario != 'complete':
            with pytest.raises(RuntimeError):
                capture.capture(**options)
            receipt = json.loads((output / 'evidence.json').read_text())
            report = json.loads((output / 'capture.json').read_text())
            assert receipt['stage'] == report['stage'] == 'failed'
            assert report['failed_stage'] == {'login': 'document', 'wrong-slide': 'ready', 'overlay': 'blocker'}[scenario]
        else:
            report = capture.capture(**options)
            receipt = json.loads((output / 'evidence.json').read_text())
            assert report['mode'] == 'fixture' and report['stage'] == 'complete'
            assert receipt['project'] == 'chromium-fixture'
            assert [(p['width'], p['height']) for p in receipt['captures']] == [(802, 452)] * 2
            assert len({p['sha256'] for p in receipt['captures']}) == 2
            # A damaged copy must fail the same upstream validator.
            import shutil
            broken = tmp_path / 'broken'
            shutil.copytree(output, broken)
            (broken / 'slide-001.png').write_bytes(b'not a png')
            command = [sys.executable, '-S', str(kit / 'scripts/check_capture_evidence.py'), str(broken),
                       '--expect', 'chromium-fixture', '--captures', 'slide-001.png,slide-002.png',
                       '--sha', options['sha'], '--run-id', options['run_id'],
                       '--run-attempt', options['attempt']]
            result = subprocess.run(command, capture_output=True, text=True, timeout=30)
            assert result.returncode == 1
    assert source.read_bytes() == before


@pytest.mark.parametrize('existing,cleanup_error', [(True, False), (False, False), (False, True)])
def test_live_failure_preserves_user_browser_and_original_error(source, tmp_path, monkeypatch, existing, cleanup_error):
    from unittest.mock import MagicMock
    page = MagicMock()
    page.url = 'https://powerpoint.cloud.microsoft/document'
    if cleanup_error:
        page.close.side_effect = OSError('cleanup details must not leak')
    context = MagicMock()
    context.pages = [page] if existing else []
    context.new_page.return_value = page
    browser = MagicMock()
    browser.contexts = [context]
    browser.version = 'test'
    pw = MagicMock()
    pw.chromium.connect_over_cdp.return_value = browser
    @contextmanager
    def session():
        yield pw
    assertion = MagicMock()
    assertion.to_have_text.side_effect = AssertionError('private URL SECRET')
    monkeypatch.setitem(sys.modules, 'playwright.sync_api', SimpleNamespace(
        sync_playwright=session, expect=lambda *_: assertion))
    monkeypatch.setattr(capture.importlib.metadata, 'version', lambda *_: 'test')
    monkeypatch.setattr(capture, 'verify_kit', lambda *_: tmp_path)
    stages = []
    monkeypatch.setattr(capture, 'seal', lambda *a, **kw: stages.append(kw['stage']))
    config = args(source, tmp_path)
    config.update(fixture=False, cdp='http://127.0.0.1:9222', url=page.url)
    with pytest.raises(RuntimeError, match='document.*AssertionError'):
        capture.capture(**config)
    report = json.loads((config['output'] / 'capture.json').read_text())
    assert report['failed_stage'] == 'document' and stages == ['failed']
    assert 'SECRET' not in json.dumps(report)
    assert page.close.call_count == (0 if existing else 1)
    browser.close.assert_not_called()
    context.close.assert_not_called()
    page.goto.assert_not_called()
    if cleanup_error:
        assert report['cleanup_error_type'] == 'OSError'
