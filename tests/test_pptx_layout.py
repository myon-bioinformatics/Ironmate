import json
from pathlib import Path
import subprocess
import sys
from zipfile import ZipFile

from xml.etree import ElementTree as ET
from tests.fixture_pptx_layout import CASES, IDS, case_rule, make_source
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches
from pptx_layout import arrange

@pytest.fixture
def deck(tmp_path):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for i in range(4):
        s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.1), Inches(.1), Inches(1 + i * .1), Inches(.5))
        s.text = f'Card {i}'
    path = tmp_path / 'source.pptx'
    prs.save(path)
    return path

def rule():
    return dict(pattern='grid', slide=1, shape_ids=[2,3,4,5], columns=2, x=1, y=1, gap_x=.2, gap_y=.3)

def test_geometry_preservation_and_preview(deck, tmp_path):
    out = tmp_path / 'out.pptx'
    before = deck.read_bytes()
    preview = arrange(deck, [rule()], dry_run=True)
    assert preview['changes'] == arrange(deck, [rule()], destination=out)['changes']
    assert deck.read_bytes() == before
    shapes = Presentation(out).slides[0].shapes
    assert [(s.left, s.top) for s in shapes] == [(Inches(1), Inches(1)), (Inches(2.5), Inches(1)), (Inches(1), Inches(1.8)), (Inches(2.5), Inches(1.8))]
    assert [(s.text,s.width,s.height) for s in shapes] == [(s.text,s.width,s.height) for s in Presentation(deck).slides[0].shapes]
    with ZipFile(deck) as a, ZipFile(out) as b:
        assert a.namelist() == b.namelist()
        assert all(a.read(n) == b.read(n) for n in a.namelist() if n != 'ppt/slides/slide1.xml')
    assert arrange(out, [rule()], dry_run=True)['change_count'] == 0

@pytest.mark.parametrize('patch', [dict(gap_x=float('nan')),dict(gap_y=-1),dict(shape_ids=[2,2]),dict(shape_ids=[99]),dict(columns=0),dict(x=99),dict(pattern='bad')])
def test_invalid_does_not_publish(deck,tmp_path,patch):
    config = rule()
    config.update(patch)
    out = tmp_path / 'bad.pptx'
    with pytest.raises(ValueError):
        arrange(deck,[config],destination=out)
    assert not out.exists()

@pytest.mark.parametrize('mode',['row','column'])
def test_sequential_preview(deck,tmp_path,mode):
    second = rule()
    second['pattern'] = mode
    out = tmp_path / 'out.pptx'
    assert arrange(deck,[rule(),second],dry_run=True)['changes'] == arrange(deck,[rule(),second],destination=out)['changes']

def test_no_overwrite(deck):
    before = deck.read_bytes()
    with pytest.raises(ValueError):
        arrange(deck,[rule()],destination=deck)
    assert deck.read_bytes() == before

def test_connected_and_rotated_shapes_rejected(deck,tmp_path):
    from pptx.enum.shapes import MSO_CONNECTOR
    prs = Presentation(deck)
    slide = prs.slides[0]
    connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,0,0,100,100)
    connector.begin_connect(slide.shapes[0],0)
    connected = tmp_path / 'connected.pptx'
    prs.save(connected)
    with pytest.raises(ValueError,match='routing'):
        arrange(connected,[rule()],dry_run=True)
    prs = Presentation(deck)
    prs.slides[0].shapes[0].rotation = 20
    rotated = tmp_path / 'rotated.pptx'
    prs.save(rotated)
    with pytest.raises(ValueError,match='unrotated'):
        arrange(rotated,[rule()],dry_run=True)


@pytest.mark.parametrize('case', CASES, ids=[case[0] for case in CASES])
def test_unequal_sizes_literal_coordinates_and_preservation(tmp_path, case):
    source, output = tmp_path / 'source.pptx', tmp_path / 'output.pptx'
    make_source(source)
    original_bytes = source.read_bytes()
    config = case_rule(case)
    preview = arrange(source, [config], dry_run=True)
    rules = tmp_path / 'rules.toml'
    rules.write_text('[[rules]]\n' + '\n'.join(
        f'{key} = {json.dumps(value)}' for key, value in config.items()) + '\n', encoding='utf-8')
    process = subprocess.run([sys.executable, str(Path(__file__).resolve().parents[1] / 'pptx_layout.py'),
                              str(source), '--rules', str(rules), '--output', str(output)],
                             capture_output=True, text=True, check=True)
    applied = json.loads(process.stdout)
    assert preview['changes'] == applied['changes']
    assert applied['change_count'] == 5
    assert source.read_bytes() == original_bytes
    before = Presentation(source).slides[0].shapes
    after = Presentation(output).slides[0].shapes
    by_id = {s.shape_id: s for s in after}
    # Explicit independent expectations, not calculated from preview or apply.
    assert [(by_id[i].left, by_id[i].top) for i in IDS] == [
        (round(x * 914400), round(y * 914400)) for x, y in case[4]]
    assert [(s.shape_id, s.name, s.text, s.width, s.height) for s in before] == [
        (s.shape_id, s.name, s.text, s.width, s.height) for s in after]
    assert (before[-1].left, before[-1].top) == (after[-1].left, after[-1].top)
    # Positive-area overlap is forbidden; touching edges for zero gap is valid.
    targets = [by_id[i] for i in IDS]
    for i, a in enumerate(targets):
        for b in targets[i + 1:]:
            assert (a.left + a.width <= b.left or b.left + b.width <= a.left
                    or a.top + a.height <= b.top or b.top + b.height <= a.top)
    # Normalizing only selected offsets must leave the complete XML unchanged:
    # covers runs/styles, shape geometry, ordering and unselected shapes.
    ns = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
          'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    with ZipFile(source) as old, ZipFile(output) as new:
        assert old.namelist() == new.namelist()
        for name in old.namelist():
            if name != 'ppt/slides/slide1.xml':
                assert old.read(name) == new.read(name)
                continue
            roots = [ET.fromstring(z.read(name)) for z in (old, new)]
            for root in roots:
                for shape in root.findall('p:cSld/p:spTree/p:sp', ns):
                    if int(shape.find('p:nvSpPr/p:cNvPr', ns).get('id')) in IDS:
                        off = shape.find('p:spPr/a:xfrm/a:off', ns)
                        off.set('x', '0')
                        off.set('y', '0')
            assert ET.canonicalize(ET.tostring(roots[0], encoding='unicode')) == ET.canonicalize(ET.tostring(roots[1], encoding='unicode'))
    assert arrange(output, [config], dry_run=True)['change_count'] == 0
