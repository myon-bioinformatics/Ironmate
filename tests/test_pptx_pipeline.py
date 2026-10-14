"""Pipeline contract tests without invoking a browser or installing Node."""
from pathlib import Path
import pytest
import pptx_pipeline

def test_pipeline_reuses_conversion_and_inspection(monkeypatch, tmp_path):
    seen = {}
    def fake_convert(source, output, *, marp):
        seen["markdown"] = Path(source).read_text(encoding="utf-8")
        return {"slides": 1}
    monkeypatch.setattr(pptx_pipeline, "convert", fake_convert)
    monkeypatch.setattr(pptx_pipeline, "inspect", lambda path: {"slides": [{"status": "unknown"}]})
    result = pptx_pipeline.run(text="# Hello\n", output=tmp_path / "slide.pptx")
    assert seen["markdown"] == "# Hello\n"
    assert result["inspection"]["slides"][0]["status"] == "unknown"

def test_mermaid_source_is_preserved(monkeypatch, tmp_path):
    fence = chr(96) * 3
    source = fence + "mermaid\nflowchart LR\n A[Start] --> B[End]\n" + fence + "\n"
    captured = []
    def fake_convert(path, output, *, marp):
        captured.append(Path(path).read_text(encoding="utf-8"))
        return {"slides": 1}
    monkeypatch.setattr(pptx_pipeline, "convert", fake_convert)
    monkeypatch.setattr(pptx_pipeline, "inspect", lambda path: {"slides": []})
    pptx_pipeline.run(text=source, output=tmp_path / "diagram.pptx")
    assert captured == [source]

def test_empty_input_is_rejected(tmp_path):
    with pytest.raises(ValueError, match="empty"):
        pptx_pipeline.run(text=" ", output=tmp_path / "slide.pptx")

def test_markdown_marp_native_enrichment_chain(monkeypatch,tmp_path):
    from pptx import Presentation
    calls=[]
    def fake_marp(source,output,*,marp):
        calls.append(Path(source).read_text())
        prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.save(output)
        return dict(slides=1,backend='marp',output=str(output))
    monkeypatch.setattr(pptx_pipeline,'convert',fake_marp)
    config=dict(pattern='flow',slide=1,x=1,y=1,width=8,height=3,gap=.4,
                items=[dict(title='A'),dict(title='B')])
    output=tmp_path/'enriched.pptx'
    result=pptx_pipeline.run(text='# Markdown\n',output=output,enrich_patterns=[config])
    assert calls==['# Markdown\n']
    assert result['conversion']['backend']=='marp'
    assert result['conversion']['output']==str(output)
    assert result['enrichment']['added_shapes']==9
    assert result['inspection']['slides'][0]['editable_runs']>0
    assert not list(tmp_path.glob('ironmate-enrich-*'))


def test_enrichment_failure_leaves_no_final_output(monkeypatch,tmp_path):
    from pptx import Presentation
    def fake_marp(source,output,*,marp):
        prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.save(output)
        return dict(slides=1)
    monkeypatch.setattr(pptx_pipeline,'convert',fake_marp)
    output=tmp_path/'bad.pptx'
    with pytest.raises(ValueError):
        pptx_pipeline.run(text='# Input',output=output,enrich_patterns=[dict(pattern='invalid')])
    assert not output.exists()
    assert not list(tmp_path.glob('ironmate-enrich-*'))

def test_cli_enrichment_resolves_csv_relative_to_toml(monkeypatch,tmp_path,capsys):
    from pptx import Presentation
    def fake_marp(source,output,*,marp):
        prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.save(output)
        return dict(slides=1,backend='marp')
    monkeypatch.setattr(pptx_pipeline,'convert',fake_marp)
    config_dir=tmp_path/'config';config_dir.mkdir()
    (config_dir/'data.csv').write_text('Key,Value\na,b\n')
    config=config_dir/'patterns.toml'
    config.write_text('[[patterns]]\npattern="table"\nslide=1\nx=1\ny=1\nwidth=8\nheight=3\ncsv="data.csv"\n')
    output=tmp_path/'out.pptx'
    pptx_pipeline.main(['--text','# Markdown','--enrich',str(config),'-o',str(output)])
    assert Presentation(output).slides[0].shapes[-1].table.cell(1,1).text=='b'
    assert '"enrichment"' in capsys.readouterr().out

@pytest.mark.parametrize('suffix,content',[('.csv','Name,Detail\nA,one\nB,two\n'),('.md','| Name | Detail |\n| --- | --- |\n| A | one |\n| B | two |\n')])
def test_direct_table_input_to_diagram(tmp_path,suffix,content):
    path=tmp_path/('data'+suffix);path.write_text(content)
    table=pptx_pipeline.table_pattern(path)
    assert table['rows']==[['Name','Detail'],['A','one'],['B','two']]
    cards=pptx_pipeline.table_pattern(path,pattern='cards',columns=2)
    assert cards['items']==[dict(title='A',body='Detail: one'),dict(title='B',body='Detail: two')]

def test_markdown_reuses_vendor_live_table_parser(tmp_path):
    path=tmp_path/'data.md'
    path.write_text('```\n| Wrong | X |\n| --- | --- |\n| nope | x |\n```\n\n| Name | Detail |\n| --- | --- |\n| A | a\\|b |\n')
    assert pptx_pipeline.table_pattern(path)['rows']==[['Name','Detail'],['A','a|b']]

def test_direct_csv_cli_needs_no_toml(monkeypatch,tmp_path,capsys):
    from pptx import Presentation
    def fake_marp(source,output,*,marp):
        prs=Presentation();prs.slides.add_slide(prs.slide_layouts[6]);prs.save(output)
        return dict(slides=1,backend='marp')
    monkeypatch.setattr(pptx_pipeline,'convert',fake_marp)
    path=tmp_path/'data.csv';path.write_text('Name,Detail\nA,one\nB,two\n')
    output=tmp_path/'out.pptx'
    pptx_pipeline.main(['--text','# Base','--table',str(path),'--table-pattern','flow','-o',str(output)])
    assert len(Presentation(output).slides[0].shapes)==9
    assert '"enrichment"' in capsys.readouterr().out

@pytest.mark.parametrize('content',['no table','| H |\n| --- |\n','```\n| H |\n| --- |\n| x |\n```'])
def test_missing_live_markdown_body_is_rejected(tmp_path,content):
    path=tmp_path/'data.md';path.write_text(content)
    with pytest.raises(ValueError):pptx_pipeline.table_pattern(path)
