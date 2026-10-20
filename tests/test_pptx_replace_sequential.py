"""Dry-run must predict cascading edits, including their exact locations."""
import json

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

import pptx_replace


def make_source(tmp_path):
    source = tmp_path / "source.pptx"
    presentation = Presentation()
    for _ in range(2):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
        for text in ("Draft Draft", " unchanged"):
            run = shape.text_frame.paragraphs[0].add_run()
            run.text = text
            run.font.color.rgb = RGBColor(255, 0, 0)
            run.font.name = "Arial"
            run.font.size = Pt(12)
            run.font.bold = False
    presentation.save(source)
    return source


def test_chained_text_reports_exact_matches_and_locations(tmp_path):
    source = make_source(tmp_path)
    before = source.read_bytes()
    output = tmp_path / "output.pptx"
    rules = [
        {"operation": "replace_text", "old": "Draft", "new": "Final", "slide": 1},
        {"operation": "replace_text", "old": "Final", "new": "Approved"},
        {"operation": "replace_text", "old": "Draft", "new": "Archived"},
    ]
    expected = [
        {"rule": rule, "slide": slide, "shape": 1, "paragraph": 1, "run": 1, "count": 2}
        for rule, slide in ((0, 1), (1, 1), (2, 2))
    ]
    preview = pptx_replace.transform(source, rules, dry_run=True, destination=output)
    assert source.read_bytes() == before
    assert set(tmp_path.iterdir()) == {source}
    actual = pptx_replace.transform(source, rules, destination=output)
    assert preview == {**actual, "dry_run": True}
    assert preview["changes"] == expected
    assert preview["change_count"] == 6
    result = Presentation(output)
    assert result.slides[0].shapes[0].text == "Approved Approved unchanged"
    assert result.slides[1].shapes[0].text == "Archived Archived unchanged"
    assert source.read_bytes() == before


@pytest.mark.parametrize("operation,values", [
    ("text_color", [{"color": "#112233"}, {"color": "#FF0000"}]),
    ("font_family", [{"name": "Meiryo"}, {"name": "Arial"}]),
    ("font_size", [{"points": 18}, {"points": 12}]),
    ("bold", [{"value": True}, {"value": False}]),
    ("slide_background", [{"color": "#112233"}, {"color": "#FFFFFF"}]),
])
def test_style_repeat_and_restore_use_previous_rule_state(tmp_path, operation, values):
    source = make_source(tmp_path)
    output = tmp_path / "output.pptx"
    # Repeating the first update is a no-op; restoring the original is an edit.
    rules = [{"operation": operation, "slide": 1, **value}
             for value in (values[0], values[0], values[1])]
    preview = pptx_replace.transform(source, rules, dry_run=True)
    actual = pptx_replace.transform(source, rules, destination=output)
    assert preview == {**actual, "dry_run": True}
    if operation == "slide_background":
        expected = [{"rule": rule, "slide": 1, "operation": operation} for rule in (0, 2)]
    else:
        expected = [{"rule": rule, "slide": 1, "shape": 1, "paragraph": 1,
                     "run": run, "count": 1} for rule in (0, 2) for run in (1, 2)]
    assert preview["changes"] == expected
    assert preview["change_count"] == len(expected)
    result = Presentation(output)
    for slide in result.slides:
        assert slide.background.fill.fore_color.rgb == RGBColor(255, 255, 255)
        for run in slide.shapes[0].text_frame.paragraphs[0].runs:
            assert run.font.color.rgb == RGBColor(255, 0, 0)
            assert run.font.name == "Arial"
            assert run.font.size == Pt(12)
            assert run.font.bold is False


def test_cli_mixed_rules_and_no_dry_run_writes(tmp_path, capsys, monkeypatch):
    source = make_source(tmp_path)
    before = source.read_bytes()
    output = tmp_path / "output.pptx"
    rules = tmp_path / "rules.toml"
    rules.write_text('''
[[rules]]
operation = "replace_text"
old = "Draft"
new = "Final"
[[rules]]
operation = "text_color"
color = "#000000"
[[rules]]
operation = "font_family"
name = "Meiryo"
[[rules]]
operation = "slide_background"
color = "#112233"
[[rules]]
operation = "font_size"
points = 18
[[rules]]
operation = "bold"
value = true
[[rules]]
operation = "replace_text"
old = "Final"
new = "Approved"
''', encoding="utf-8")
    arguments = [str(source), "--rules", str(rules), "--output", str(output)]
    presentation_type = type(Presentation(source))

    def no_save(*args, **kwargs):
        pytest.fail("dry-run must never call save")

    with monkeypatch.context() as patch:
        patch.setattr(presentation_type, "save", no_save)
        pptx_replace.main([*arguments, "--dry-run"])
        preview = json.loads(capsys.readouterr().out)
        # Even an output aliasing the source must leave the original untouched.
        pptx_replace.main([str(source), "--rules", str(rules), "--output", str(source), "--dry-run"])
        assert json.loads(capsys.readouterr().out) == preview
    assert source.read_bytes() == before
    assert set(tmp_path.iterdir()) == {source, rules}
    pptx_replace.main(arguments)
    actual = json.loads(capsys.readouterr().out)
    assert preview == {**actual, "dry_run": True}
    assert preview["change_count"] == 26
    for slide in Presentation(output).slides:
        assert slide.shapes[0].text == "Approved Approved unchanged"
        assert slide.background.fill.fore_color.rgb == RGBColor(17, 34, 51)
        for run in slide.shapes[0].text_frame.paragraphs[0].runs:
            assert run.font.color.rgb == RGBColor(0, 0, 0)
            assert run.font.name == "Meiryo"
            assert run.font.size == Pt(18)
            assert run.font.bold is True
    assert source.read_bytes() == before


@pytest.mark.parametrize("dry_run", [True, False])
@pytest.mark.parametrize("rule,message", [
    ({"operation": "arrow"}, "unsupported operation"),
    ({"operation": "replace_text", "old": "", "new": "x"}, "nonempty old"),
    ({"operation": "text_color", "color": "red"}, "#RRGGBB"),
    ({"operation": "slide_background", "color": "#GGGGGG"}, "#RRGGBB"),
    ({"operation": "font_family", "name": " "}, "nonempty name"),
    ({"operation": "font_size", "points": True}, "numeric points"),
    ({"operation": "bold", "value": "true"}, "boolean value"),
    ({"operation": "bold", "value": True, "slide": 3}, "slide must be"),
])
def test_invalid_later_rule_fails_without_writes(tmp_path, dry_run, rule, message):
    source = make_source(tmp_path)
    before = source.read_bytes()
    with pytest.raises(ValueError, match=message):
        pptx_replace.transform(source, [
            {"operation": "replace_text", "old": "Draft", "new": "Final"}, rule,
        ], dry_run=dry_run, destination=tmp_path / "output.pptx")
    assert source.read_bytes() == before
    assert set(tmp_path.iterdir()) == {source}


@pytest.mark.parametrize("dry_run", [True, False])
@pytest.mark.parametrize("operation", ["text_color", "slide_background"])
def test_inherited_colors_still_fail_after_prior_edit(tmp_path, dry_run, operation):
    source = tmp_path / "source.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Draft"
    presentation.save(source)
    before = source.read_bytes()
    with pytest.raises(ValueError, match="explicit RGB|explicit solid RGB|inherited slide background"):
        pptx_replace.transform(source, [
            {"operation": "replace_text", "old": "Draft", "new": "Final"},
            {"operation": operation, "color": "#112233"},
        ], dry_run=dry_run, destination=tmp_path / "output.pptx")
    assert source.read_bytes() == before
    assert set(tmp_path.iterdir()) == {source}
