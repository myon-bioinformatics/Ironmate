"""PPTX-first styling contracts and explicit Open XML fallback boundaries."""
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import pytest
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from pptx_replace import transform

NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


def sample(path, *, explicit_background=True, explicit_text=True):
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    if explicit_background:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(25, 50, 75)
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    run = shape.text_frame.paragraphs[0].add_run()
    run.text = "Heading"
    if explicit_text:
        run.font.color.rgb = RGBColor(200, 100, 50)
    prs.save(path)


def test_plain_style_batch_is_repeatable_and_preserves_other_parts(tmp_path):
    original = tmp_path / "original.pptx"
    first = tmp_path / "first.pptx"
    second = tmp_path / "second.pptx"
    sample(original)
    rules = [
        {"operation": "slide_background", "color": "#FFFFFF"},
        {"operation": "text_color", "color": "#000000"},
        {"operation": "font_family", "name": "Meiryo"},
        {"operation": "font_size", "points": 20},
        {"operation": "bold", "value": True},
    ]
    assert transform(original, rules, dry_run=True)["change_count"] == 5
    assert transform(original, rules, destination=first)["change_count"] == 5
    assert transform(first, rules, dry_run=True)["change_count"] == 0
    assert transform(first, rules, destination=second)["change_count"] == 0
    prs = Presentation(second)
    slide = prs.slides[0]
    run = slide.shapes[0].text_frame.paragraphs[0].runs[0]
    assert slide.background.fill.fore_color.rgb == RGBColor(255, 255, 255)
    assert run.text == "Heading"
    assert run.font.color.rgb == RGBColor(0, 0, 0)
    assert run.font.name == "Meiryo"
    assert run.font.size.pt == 20
    assert run.font.bold is True
    with ZipFile(original) as before, ZipFile(first) as after:
        names = set(before.namelist())
        assert names == set(after.namelist())
        assert all(before.read(name) == after.read(name)
                   for name in names if name != "ppt/slides/slide1.xml")
        xml = ET.fromstring(after.read("ppt/slides/slide1.xml"))
        assert xml.find(".//a:rPr/a:solidFill/a:srgbClr", NS).attrib["val"] == "000000"


@pytest.mark.parametrize("explicit_background,explicit_text,operation,expected", [
    (False, True, "slide_background", "only explicit solid RGB slide backgrounds"),
    (True, False, "text_color", "explicit RGB"),
])
def test_unverified_inherited_styles_fail_closed(tmp_path, explicit_background,
                                                   explicit_text, operation, expected):
    source = tmp_path / "source.pptx"
    target = tmp_path / "target.pptx"
    sample(source, explicit_background=explicit_background, explicit_text=explicit_text)
    rule = {"operation": operation, "color": "#FFFFFF" if operation == "slide_background" else "#000000"}
    with pytest.raises(ValueError, match=expected):
        transform(source, [rule], destination=target)
    assert not target.exists()
    assert Presentation(source).slides[0].shapes[0].text == "Heading"
