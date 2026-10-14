from zipfile import ZipFile
import pytest
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches
from pptx_enrich import enrich, main

@pytest.fixture
def deck(tmp_path):
    prs=Presentation()
    for _ in range(2):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(.2),Inches(.2))
        s.text='Keep me'
    source=tmp_path/'input.pptx'
    prs.save(source)
    return source

def spec(mode):
    result=dict(pattern=mode,slide=1,x=.5,y=1,width=9,height=4,gap=.25,
                items=[dict(title=f'Item {i}',body='Line 1\nLine 2',label='1–3') for i in range(3)])
    if mode=='table':
        result['rows']=[['A','B'],['a','b'],['c','d']]
    return result

@pytest.mark.parametrize('mode',['cards','regions','flow','agenda','table'])
def test_native_objects_preview_and_preservation(deck,tmp_path,mode):
    before=deck.read_bytes()
    out=tmp_path/'out.pptx'
    planned=enrich(deck,[spec(mode)],dry_run=True)
    actual=enrich(deck,[spec(mode)],destination=out)
    assert planned['patterns']==actual['patterns']
    assert deck.read_bytes()==before
    slide=Presentation(out).slides[0]
    assert slide.shapes[0].text=='Keep me'
    assert all(s.left>=0 and s.top>=0 and s.left+s.width<=Inches(10) and s.top+s.height<=Inches(7.5) for s in slide.shapes)
    assert not any(s.shape_type==13 for s in slide.shapes)
    assert len({s.shape_id for s in slide.shapes})==len(slide.shapes)
    if mode=='table':
        table=slide.shapes[-1].table
        assert table.cell(1,1).text=='b'
        table.cell(1,1).text='edited'
    else:
        texts=[s.text for s in slide.shapes if s.has_text_frame]
        assert any('Item 0' in t for t in texts)
        arrows = [s for s in slide.shapes if s.name.endswith('-arrow')]
        assert len(arrows)==(2 if mode=='flow' else 0)
        for arrow in arrows:
            assert arrow.height == 0
            assert arrow._element.xpath('.//a:tailEnd')[0].get('type') == 'triangle'
    with ZipFile(deck) as a,ZipFile(out) as b:
        assert a.namelist()==b.namelist()
        assert all(a.read(n)==b.read(n) for n in a.namelist() if n!='ppt/slides/slide1.xml')

@pytest.mark.parametrize('patch',[dict(x=float('nan')),dict(gap=-.2),dict(width=30),dict(columns=0),dict(items=[]),dict(pattern='bad'),dict(font_size=3),dict(accent='red'),dict(slide=99)])
def test_invalid_no_publication(deck,tmp_path,patch):
    config=spec('cards');config.update(patch)
    out=tmp_path/'bad.pptx'
    with pytest.raises(ValueError):enrich(deck,[config],destination=out)
    assert not out.exists()

def test_multiple_patterns_same_slide(deck,tmp_path):
    out=tmp_path/'multi.pptx'
    a=spec('cards');a['height']=1
    b=spec('agenda');b['height']=3;b['y']=3
    result=enrich(deck,[a,b],destination=out)
    assert len(Presentation(out).slides[0].shapes)==1+result['added_shapes']

def test_no_overwrite_and_later_error(deck,tmp_path):
    before=deck.read_bytes()
    with pytest.raises(ValueError):enrich(deck,[spec('cards')],destination=deck)
    bad=spec('table');bad['rows']=[['a'],['b','c']]
    out=tmp_path/'bad.pptx'
    with pytest.raises(ValueError):enrich(deck,[spec('cards'),bad],destination=out)
    assert not out.exists() and deck.read_bytes()==before

def test_cli(deck,tmp_path,capsys):
    config=tmp_path/'spec.toml'
    config.write_text('[[patterns]]\npattern="agenda"\nslide=1\nx=1\ny=1\nwidth=8\nheight=3\n[[patterns.items]]\ntitle="Overview"\nbody="Description"\n')
    main([str(deck),'--patterns',str(config),'--dry-run'])
    assert '"added_shapes": 2' in capsys.readouterr().out

def test_shipped_example_roundtrip(tmp_path):
    import tomllib
    from pathlib import Path
    prs=Presentation()
    for _ in range(4):prs.slides.add_slide(prs.slide_layouts[6])
    source=tmp_path/'blank.pptx';prs.save(source)
    with (Path(__file__).parents[1]/'examples/pptx_enrich.toml').open('rb') as stream:
        patterns=tomllib.load(stream)['patterns']
    for config in patterns:
        if 'csv' in config:config['csv']=str(Path(__file__).parents[1]/'examples'/config['csv'])
    out=tmp_path/'examples.pptx'
    result=enrich(source,patterns,destination=out)
    assert [p['pattern'] for p in result['patterns']]==['cards','regions','flow','table']
    assert all(len(slide.shapes)>0 for slide in Presentation(out).slides)

def test_csv_values_and_explicit_table_colors(deck,tmp_path):
    path=tmp_path/'comparison.csv'
    path.write_text('\ufeffCriterion,Value\n"quoted, field",日本語\nNext,42\n',encoding='utf-8')
    config=spec('table');del config['rows']
    config.update(csv=str(path),header_fill='#202633',row_fills=['#FFFFFF','#EAF1F8'])
    out=tmp_path/'csv.pptx'
    enrich(deck,[config],destination=out)
    table=Presentation(out).slides[0].shapes[-1].table
    assert table.cell(1,0).text=='quoted, field'
    assert table.cell(1,1).text=='日本語'
    assert str(table.cell(0,0).fill.fore_color.rgb)=='202633'
    assert str(table.cell(1,0).fill.fore_color.rgb)=='FFFFFF'
    assert str(table.cell(2,0).fill.fore_color.rgb)=='EAF1F8'
    assert len(table.cell(0,0)._tc.xpath('./a:tcPr/a:lnL'))==1

@pytest.mark.parametrize('patch',[dict(row_fills=[]),dict(row_fills=['wrong']),dict(header_fill='bad')])
def test_invalid_table_colors(deck,tmp_path,patch):
    config=spec('table');config.update(patch)
    with pytest.raises(ValueError):enrich(deck,[config],dry_run=True)

def test_default_table_is_neutral_grayscale(deck,tmp_path):
    out=tmp_path/'gray.pptx'
    enrich(deck,[spec('table')],destination=out)
    table=Presentation(out).slides[0].shapes[-1].table
    assert str(table.cell(0,0).fill.fore_color.rgb)=='222222'
    assert str(table.cell(1,0).fill.fore_color.rgb)=='FFFFFF'
    assert str(table.cell(2,0).fill.fore_color.rgb)=='F2F2F2'
