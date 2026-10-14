"""Synthetic layout acceptance fixtures; no reference/user material is included.

Run from the repository root: python -m tests.fixture_pptx_layout OUTPUT_DIR
"""
import argparse
import hashlib
import json
import subprocess
import tempfile

from PIL import Image, ImageChops
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

from pptx_layout import arrange

# Literal acceptance coordinates in rule ID order. Do not derive these with the
# layout algorithm: they are the independent oracle for tests and reference decks.
CASES = [
    ('row-zero', 'row', 0, .9, [(.75, 1), (2.75, 1), (4.75, 1), (6.75, 1), (8.75, 1)]),
    ('row-gap', 'row', .4, .9, [(.75, 1), (3.15, 1), (5.55, 1), (7.95, 1), (10.35, 1)]),
    ('column-zero', 'column', .9, 0, [(.75, 1), (.75, 2.25), (.75, 3.5), (.75, 4.75), (.75, 6)]),
    ('column-gap', 'column', .9, .3, [(.75, 1), (.75, 2.55), (.75, 4.1), (.75, 5.65), (.75, 7.2)]),
    ('grid-zero', 'grid', 0, 0, [(.75, 1), (2.75, 1), (4.75, 1), (.75, 2.25), (2.75, 2.25)]),
    ('grid-gap', 'grid', .2, .3, [(.75, 1), (2.95, 1), (5.15, 1), (.75, 2.55), (2.95, 2.55)]),
]
IDS = [6, 2, 5, 3, 4]  # Deliberately different from document order.


def make_source(path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(14), Inches(10)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for i, (width, height) in enumerate([(1, .5), (2, .75), (1.5, 1.25), (.75, 1), (1.25, .6)]):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(10 + i * .15),
                                       Inches(8 + i * .1), Inches(width), Inches(height))
        shape.name = f'Synthetic target {i + 1}'
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(240 - i * 15, 240 - i * 15, 240 - i * 15)
        shape.line.color.rgb = RGBColor(0, 0, 0)
        frame = shape.text_frame
        frame.word_wrap = False
        for text, bold in [('Box ', True), (str(i + 1), False)]:
            run = frame.paragraphs[0].add_run()
            run.text = text
            run.font.name, run.font.size = 'DejaVu Sans', Pt(10)
            run.font.bold = bold
            run.font.color.rgb = RGBColor(0, 0, 0)
    note = slide.shapes.add_textbox(Inches(.75), Inches(.15), Inches(8), Inches(.5))
    note.text = 'Synthetic layout fixture: unchanged label'
    note.text_frame.paragraphs[0].runs[0].font.name = 'DejaVu Sans'
    note.text_frame.paragraphs[0].runs[0].font.size = Pt(16)
    prs.save(path)


def case_rule(case):
    name, mode, gx, gy, expected = case
    return dict(pattern=mode, slide=1, shape_ids=IDS, columns=3,
                x=.75, y=1, gap_x=gx, gap_y=gy)


def generate(directory):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=False)
    source = directory / 'source.pptx'
    make_source(source)
    manifest = {'scope': 'Synthetic native shapes only; PowerPoint Web unverified', 'cases': []}
    for case in CASES:
        name, mode, gx, gy, expected = case
        rule = case_rule(case)
        rules = directory / f'{name}.toml'
        rules.write_text('[[rules]]\n' + '\n'.join(
            f'{k} = {json.dumps(v)}' for k, v in rule.items()) + '\n', encoding='utf-8')
        output = directory / f'{name}.pptx'
        result = arrange(source, [rule], destination=output)
        # Independent reference: python-pptx setters and literal coordinates,
        # never arrange(), its preview, or coordinates read from its output.
        reference = Presentation(source)
        by_id = {s.shape_id: s for s in reference.slides[0].shapes}
        for shape_id, (left, top) in zip(IDS, expected):
            by_id[shape_id].left = round(left * 914400)
            by_id[shape_id].top = round(top * 914400)
        reference.save(directory / f'{name}-expected.pptx')
        manifest['cases'].append({'name': name, 'rule': rule,
                                  'expected_inches_in_id_order': expected, 'result': result})
    manifest['files'] = {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                         for p in sorted(directory.iterdir())}
    (directory / 'manifest.json').write_text(json.dumps(manifest, indent=2) + '\n', encoding='utf-8')
    return manifest


def render(directory):
    """Optional local acceptance rendering; never used by default pytest/CI.

    Requires already installed soffice, pdfinfo, pdftoppm, pdftotext and Pillow.
    Failure/missing tools raise, rather than claiming visual success.
    """
    directory = Path(directory).resolve()
    pdf_dir, png_dir = directory / 'pdf', directory / 'png'
    pdf_dir.mkdir(exist_ok=False)
    png_dir.mkdir(exist_ok=False)
    versions = {}
    for tool, flag in [('soffice', '--version'), ('pdftoppm', '-v'), ('pdftotext', '-v')]:
        process = subprocess.run([tool, flag], capture_output=True, text=True, check=True, timeout=30)
        versions[tool] = (process.stdout + process.stderr).strip()
    with tempfile.TemporaryDirectory(prefix='layout-render-') as profile:
        process = subprocess.run(
            ['soffice', '--headless', f'-env:UserInstallation={Path(profile).as_uri()}',
             '--convert-to', 'pdf', '--outdir', str(pdf_dir),
             *[str(p) for p in sorted(directory.glob('*.pptx'))]],
            capture_output=True, text=True, check=True, timeout=120)
        (directory / 'renderer.log').write_text(process.stdout + process.stderr, encoding='utf-8')
    for pptx in sorted(directory.glob('*.pptx')):
        pdf = pdf_dir / (pptx.stem + '.pdf')
        info = subprocess.run(['pdfinfo', str(pdf)], capture_output=True, text=True, check=True, timeout=30)
        pages = [line.split(':', 1)[1].strip() for line in info.stdout.splitlines() if line.startswith('Pages:')]
        assert pages == ['1'], (pptx.name, pages)
        subprocess.run(['pdftoppm', '-scale-to', '1120', '-singlefile', '-png', str(pdf),
                        str(png_dir / pptx.stem)], capture_output=True, check=True, timeout=30)
        subprocess.run(['pdftotext', str(pdf), str(pdf.with_suffix('.txt'))], check=True, timeout=30)
    report = {'renderer_versions': versions, 'powerpoint_web': 'unverified',
              'comparison': 'actual vs independently positioned reference, same renderer/fonts',
              'cases': []}
    for case in CASES:
        name = case[0]
        actual = Image.open(png_dir / f'{name}.png').convert('RGB')
        expected = Image.open(png_dir / f'{name}-expected.png').convert('RGB')
        assert actual.size == expected.size == (1120, 800)
        diff = ImageChops.difference(actual, expected)
        red, green, blue = diff.split()
        histogram = ImageChops.lighter(ImageChops.lighter(red, green), blue).histogram()
        changed = sum(histogram[1:])
        diff.save(png_dir / f'{name}-diff.png')
        text = (pdf_dir / f'{name}.txt').read_text(encoding='utf-8')
        assert all(f'Box {i}' in text for i in range(1, 6)), text
        assert 'Synthetic layout fixture: unchanged label' in text
        report['cases'].append({'name': name, 'changed_pixels': changed,
                                'diff_bbox': diff.getbbox(), 'image_size': actual.size,
                                'all_labels_extracted': True})
        assert changed == 0, name
    # Contact sheets only compose real renderer output. Columns are documented
    # in the report: original, independent reference, actual. Rows: zero, gap.
    for mode in ('row', 'column', 'grid'):
        sheet = Image.new('RGB', (1680, 800), 'white')
        for row, suffix in enumerate(('zero', 'gap')):
            name = f'{mode}-{suffix}'
            for col, stem in enumerate(('source', f'{name}-expected', name)):
                with Image.open(png_dir / f'{stem}.png') as img:
                    sheet.paste(img.resize((560, 400)), (col * 560, row * 400))
        sheet.save(directory / f'{mode}-comparison.png')
    report['contact_sheet'] = {'columns': ['source', 'independent expected', 'actual'],
                               'rows': ['zero gap', 'changed gap']}
    report['files'] = {str(p.relative_to(directory)): hashlib.sha256(p.read_bytes()).hexdigest()
                       for p in sorted(directory.rglob('*')) if p.is_file()}
    (directory / 'render-report.json').write_text(json.dumps(report, indent=2) + '\n', encoding='utf-8')
    return report


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('directory', type=Path)
    parser.add_argument('--render', action='store_true')
    args = parser.parse_args()
    generate(args.directory)
    if args.render:
        render(args.directory)
