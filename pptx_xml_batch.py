"""Batch-edit explicit RGB native shape fills across ordered PPTX slides.

Resolve presentation order through python-pptx slide parts, not slideN.xml guesses.
The original archive is never overwritten. All targets validate before writing.
"""
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from pptx import Presentation

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def edit_shape_fills(source, destination, rules):
    """Rules: list of {slide: 1-based display index, shape_name: str, fill_rgb: #RRGGBB}."""
    source, destination = Path(source), Path(destination)
    if source.resolve() == destination.resolve() or destination.exists():
        raise ValueError("destination must be a new, distinct file")
    if not isinstance(rules, list) or not rules:
        raise ValueError("nonempty rules list required")
    presentation = Presentation(source)
    parts = [str(s.part.partname).lstrip("/") for s in presentation.slides]
    parsed = {}
    targets = set()
    with ZipFile(source) as package:
        entries = package.infolist()
        names = [entry.filename for entry in entries]
        if len(names) != len(set(names)) or any(part not in names for part in parts):
            raise ValueError("invalid slide package members")
        for rule in rules:
            if not isinstance(rule, dict) or set(rule) != {"slide", "shape_name", "fill_rgb"}:
                raise ValueError("rule requires slide, shape_name, fill_rgb")
            index, shape_name, rgb = rule["slide"], rule["shape_name"], rule["fill_rgb"]
            if type(index) is not int or not 1 <= index <= len(parts):
                raise ValueError("slide index out of range")
            if not isinstance(shape_name, str) or not shape_name.strip():
                raise ValueError("nonempty shape_name required")
            if not isinstance(rgb, str) or len(rgb) != 7 or not rgb.startswith("#") or any(
                    ch not in "0123456789abcdefABCDEF" for ch in rgb[1:]):
                raise ValueError("fill_rgb must be #RRGGBB")
            key = (index, shape_name)
            if key in targets:
                raise ValueError("duplicate edit target")
            targets.add(key)
            member = parts[index - 1]
            if member not in parsed:
                parsed[member] = ET.fromstring(package.read(member))
            root = parsed[member]
            shapes = [
                shape for shape in root.findall(".//" + P + "sp")
                if shape.find(P + "nvSpPr/" + P + "cNvPr") is not None
                and shape.find(P + "nvSpPr/" + P + "cNvPr").get("name") == shape_name
            ]
            if len(shapes) != 1:
                raise ValueError("expected exactly one named native shape per slide")
            props = shapes[0].find(P + "spPr")
            fill = props.find(A + "solidFill") if props is not None else None
            if fill is None or len(fill) != 1 or fill[0].tag != A + "srgbClr":
                raise ValueError("explicit RGB shape fill required")
            fill[0].set("val", rgb[1:].upper())
        replacements = {
            member: ET.tostring(root, encoding="utf-8", xml_declaration=True)
            for member, root in parsed.items()
        }
        try:
            with ZipFile(destination, "x") as output:
                for entry in entries:
                    output.writestr(entry, replacements.get(entry.filename, package.read(entry.filename)))
            check = Presentation(destination)
            if len(check.slides) != len(parts):
                raise ValueError("slide count changed")
        except BaseException:
            destination.unlink(missing_ok=True)
            raise
    return {"changed_members": sorted(replacements), "changed_slides": sorted({key[0] for key in targets})}
