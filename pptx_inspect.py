"""Inspect generated PPTX text editability without assuming raster content is editable."""
import argparse
import json
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from pptx import Presentation

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def inspect(path):
    path = Path(path)
    presentation = Presentation(path)
    slides = []
    with ZipFile(path) as package:
        for number, slide in enumerate(presentation.slides, 1):
            root = ET.fromstring(package.read(f"ppt/slides/slide{number}.xml"))
            runs = sum(len(paragraph.runs) for shape in slide.shapes
                       if shape.has_text_frame for paragraph in shape.text_frame.paragraphs)
            pictures = len(root.findall(f".//{P}pic"))
            texts = [node.text or "" for node in root.findall(f".//{A}t")]
            # A slide may have mixed raster and native text. Do not infer
            # all-or-nothing editability from the presence of one text run.
            status = "editable" if runs and not pictures else "unknown" if runs else "unsupported"
            slides.append({"slide": number, "status": status, "editable_runs": runs,
                           "pictures": pictures, "text_nodes": len(texts)})
    return {"source": str(path), "slides": slides}


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("-o", "--output", type=Path)
    args = parser.parse_args(argv)
    report = json.dumps(inspect(args.source), ensure_ascii=False, indent=2) + "\n"
    if args.output:
        with args.output.open("x", encoding="utf-8") as stream:
            stream.write(report)
    else:
        print(report, end="")


if __name__ == "__main__":
    main()
